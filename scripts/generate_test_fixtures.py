# /// script
# requires-python = ">=3.12"
# dependencies = ["fpdf2", "pillow", "python-docx", "openpyxl"]
# ///
"""Generate the binary and document fixtures under server/tests/fixtures (issue E0.8).

Run from the repo root:   uv run scripts/generate_test_fixtures.py
Only SHA256SUMS:          uv run scripts/generate_test_fixtures.py --sums
Only the Outlook .msg:    python scripts/generate_test_fixtures.py --msg   (standard library only)

All content is synthetic (fictional company, example.com addresses, no phone numbers). The committed
bytes ARE the fixtures: regenerate on purpose, then commit the files together with SHA256SUMS. Dates are
fixed so reruns are close to byte-identical, but zip-based files (docx, xlsx) may still differ.
Not generated here: sow/sow_rfq_nav9.json and README.md (authored by hand).
"""

import argparse
import datetime as dt
import email.policy
import email.utils
import gzip
import hashlib
import io
import json
import shutil
import struct
import tarfile
import textwrap
import zipfile
from email.message import EmailMessage
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
FIX = REPO / "server/tests/fixtures"
WHEN = dt.datetime(2026, 10, 1, 9, 30, tzinfo=dt.UTC)
ZIP_TIME = (2026, 10, 1, 9, 30, 0)
NOT_FIXTURES = {"README.md", "SHA256SUMS", "__init__.py", "check_fixture_readme.py"}  # not checksummed

RFQ_PAGE_1 = [
    "Request for Quotation ADX-2026-014",
    "Alpha Dynamics Ltd requests a quotation for 2,500 units of the NAV-9 flight controller "
    "board, delivered to the Austin, TX facility by 2026-12-01.",
    "The assembly must meet IPC-A-610 Class 3 and AS9100D. Pricing is live costing with FX "
    "as of the quotation date and a scrap rate of 1.5 percent.",
    "Contact: sourcing@example.com. This document is synthetic test data.",
]
RFQ_PAGE_2 = [
    "Commercial terms",
    "Payment is net 30 days from invoice. Quotes are valid for 14 days. A single-source part "
    "(U-210) must be flagged for sourcing review before the quotation is released.",
]
BOM_HEADER = ["CPN", "MPN", "Qty", "Supplier", "Unit price USD"]
BOM_ROWS = [
    ["C-100", "GRM188R71H104KA93", 120, "DigiKey", 0.012],
    ["U-210", "STM32H743VIT6", 1, "Mouser", 14.75],
    ["R-330", "RC0603FR-0710KL", 85, "DigiKey", 0.004],
    ["C-415", "CL10A106KP8NNNC", 40, "Farnell", 0.031],
    ["L-050", "SRN4018-4R7M", 6, "Mouser", 0.42],
]

TEXT_FILES: dict[str, str] = {
    "sample.md": (
        "# NAV-9 RFQ summary\n\n"
        "**Customer:** Alpha Dynamics Ltd (synthetic)\n\n"
        "## Scope\n\n- 2,500 units of the NAV-9 flight controller board\n"
        "- IPC-A-610 Class 3, AS9100D\n\n"
        "## Open points\n\n1. U-210 is single-source.\n2. FX as of the quotation date.\n"
    ),
    "sample.txt": (
        "RFQ ADX-2026-014 (synthetic)\n"
        "Delivery: Austin, TX facility, 2026-12-01.\n"
        "Scrap rate: 1.5 percent. Payment: net 30.\n"
    ),
    "sample.csv": "CPN,MPN,Qty,Supplier\n"
    + "".join(f"{r[0]},{r[1]},{r[2]},{r[3]}\n" for r in BOM_ROWS),
    "sample.json": json.dumps(
        {"rule": "RULE-BOM", "policy": "flag single-source parts", "scrap_rate": 0.015,
         "approved_suppliers": ["DigiKey", "Mouser", "Farnell"]},
        indent=2,
    )
    + "\n",
    "sample.yaml": (
        "rule: RULE-OFFER\n"
        "policy: convert offers with approved FX\n"
        "currency: USD\n"
        "validity_days: 14\n"
        "stop_conditions:\n  - expired offer\n  - missing FX rate\n"
    ),
    "sample.xml": (
        '<?xml version="1.0" encoding="UTF-8"?>\n<bom project="NAV-9">\n'
        + "".join(f'  <line cpn="{r[0]}" mpn="{r[1]}" qty="{r[2]}"/>\n' for r in BOM_ROWS)
        + "</bom>\n"
    ),
    "openapi.yaml": (
        "openapi: 3.0.3\ninfo:\n  title: Quote API (synthetic)\n  version: 1.0.0\n"
        "paths:\n  /quotes:\n    get:\n      summary: List quotes\n      responses:\n"
        "        '200':\n          description: A list of quotes\n"
        "    post:\n      summary: Create a quote\n      responses:\n"
        "        '201':\n          description: Created\n"
    ),
    "fastmcp_spec.json": json.dumps(
        {"name": "bom-scrub", "version": "0.1.0",
         "tools": [{"name": "scrub_bom", "description": "Normalize a BOM and flag single-source parts",
                    "input_schema": {"type": "object", "required": ["bom_ref"],
                                     "properties": {"bom_ref": {"type": "string"}}}}]},
        indent=2,
    )
    + "\n",
    "sample.py": (
        '"""Extended cost of a BOM line (synthetic sample)."""\n\n\n'
        "def extended_cost(qty: int, unit_price: float, scrap_rate: float = 0.015) -> float:\n"
        '    """Cost including scrap."""\n'
        "    return round(qty * unit_price * (1 + scrap_rate), 4)\n\n\n"
        "class Line:\n    def __init__(self, cpn: str, qty: int) -> None:\n"
        "        self.cpn = cpn\n        self.qty = qty\n"
    ),
    "sample.sql": (
        "CREATE TABLE bom_line (\n  cpn TEXT PRIMARY KEY,\n  mpn TEXT NOT NULL,\n"
        "  qty INTEGER NOT NULL CHECK (qty > 0)\n);\n\n"
        "SELECT cpn, qty FROM bom_line WHERE qty > 10 ORDER BY cpn;\n"
    ),
    "sample.js": (
        "// Extended cost of a BOM line (synthetic sample)\n"
        "function extendedCost(qty, unitPrice, scrapRate = 0.015) {\n"
        "  return Math.round(qty * unitPrice * (1 + scrapRate) * 1e4) / 1e4;\n}\n\n"
        "module.exports = { extendedCost };\n"
    ),
    "sample.go": (
        "package bom\n\n// ExtendedCost returns the cost of a line including scrap.\n"
        "func ExtendedCost(qty int, unitPrice, scrapRate float64) float64 {\n"
        "\treturn float64(qty) * unitPrice * (1 + scrapRate)\n}\n"
    ),
    "sample.cpp": (
        "// Extended cost of a BOM line (synthetic sample)\n#include <cstdint>\n\n"
        "double extended_cost(std::int32_t qty, double unit_price, double scrap_rate) {\n"
        "    return qty * unit_price * (1.0 + scrap_rate);\n}\n"
    ),
}


def write(path: Path, data: bytes | str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, str):
        path.write_text(data, encoding="utf-8", newline="\n")
    else:
        path.write_bytes(data)


def render_page(lines: list[str]):  # type: ignore[no-untyped-def]
    """A synthetic 'scanned' page: grayscale text on white, 100 dpi A4. Deterministic."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("L", (827, 1169), 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=24)
    y = 70
    for line in lines:
        for chunk in textwrap.wrap(line, width=52) or [""]:
            draw.text((60, y), chunk, fill=0, font=font)
            y += 38
        y += 24
    return image


def png_bytes(image) -> bytes:  # type: ignore[no-untyped-def]
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def new_pdf():  # type: ignore[no-untyped-def]
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.compress = False  # keeps text readable for the PII scan
    pdf.set_creation_date(WHEN)
    pdf.set_title("Synthetic fixture")
    pdf.set_author("SBAM test fixtures")
    return pdf


def add_text_page(pdf, lines: list[str]) -> None:  # type: ignore[no-untyped-def]
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)


def add_image_page(pdf, png: bytes) -> None:  # type: ignore[no-untyped-def]
    pdf.add_page()
    pdf.image(io.BytesIO(png), x=0, y=0, w=210, h=297)


def make_pdfs(images_dir: Path) -> tuple[bytes, bytes]:
    page1, page2 = png_bytes(render_page(RFQ_PAGE_1)), png_bytes(render_page(RFQ_PAGE_2))
    text = new_pdf()
    add_text_page(text, RFQ_PAGE_1)
    add_text_page(text, RFQ_PAGE_2)
    write(images_dir / "text.pdf", bytes(text.output()))

    full = new_pdf()
    add_image_page(full, page1)
    add_image_page(full, page2)
    full_bytes = bytes(full.output())
    write(images_dir / "scanned_full.pdf", full_bytes)

    mixed = new_pdf()  # edge case: one real text page, one image-only page
    add_text_page(mixed, RFQ_PAGE_1)
    add_image_page(mixed, page2)
    write(images_dir / "scanned_mixed.pdf", bytes(mixed.output()))
    return page1, full_bytes


def make_images(directory: Path, png: bytes) -> None:
    from PIL import Image

    write(directory / "scan.png", png)
    image = Image.open(io.BytesIO(png))
    jpeg = io.BytesIO()
    image.convert("RGB").save(jpeg, format="JPEG", quality=80)
    write(directory / "photo.jpg", jpeg.getvalue())
    tiff = io.BytesIO()
    image.save(tiff, format="TIFF", compression="tiff_lzw")
    write(directory / "scan.tiff", tiff.getvalue())
    write(directory / "fake_pdf_is_png.pdf", png)  # misleading extension: PNG magic bytes


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("NAV-9 RFQ summary", level=1)
    for line in RFQ_PAGE_1[1:]:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=3)
    for cell, title in zip(table.rows[0].cells, ["CPN", "MPN", "Qty"], strict=True):
        cell.text = title
    for row in BOM_ROWS[:3]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row[:3], strict=True):
            cell.text = str(value)
    props = document.core_properties
    props.author = props.last_modified_by = "SBAM test fixtures"
    props.created = props.modified = WHEN.replace(tzinfo=None)
    props.title = "Synthetic fixture"
    document.save(path)


def make_xlsx(path: Path) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    bom = book.active
    bom.title = "BOM"
    bom["A1"] = "NAV-9 bill of materials (synthetic)"
    bom.merge_cells("A1:E1")
    bom.append(BOM_HEADER)
    for row in BOM_ROWS:
        bom.append(row)
    suppliers = book.create_sheet("Suppliers")
    suppliers.append(["Supplier", "Terms", "Minimum order USD"])
    suppliers.append(["DigiKey", "net 30", 250])
    suppliers.append(["Mouser", "net 45", 200])
    notes = book.create_sheet("Notes")
    notes["A1"] = "Review notes"
    notes.merge_cells("A1:B2")
    notes["A3"] = "U-210 is single-source; flag for sourcing review."
    book.properties.creator = book.properties.lastModifiedBy = "SBAM test fixtures"
    book.properties.created = book.properties.modified = WHEN.replace(tzinfo=None)
    book.save(path)


def make_eml(path: Path, attachment: bytes) -> None:
    message = EmailMessage(policy=email.policy.default)
    message["From"] = "rfq-desk@example.com"
    message["To"] = "sourcing@example.com"
    message["Subject"] = "RFQ ADX-2026-014 with scanned drawing"
    message["Date"] = email.utils.format_datetime(WHEN)
    message["Message-ID"] = "<fixture-eml-1@example.com>"
    message.set_content("Hello,\n\nplease find the scanned RFQ attached (synthetic test data).\n")
    message.add_attachment(attachment, maintype="application", subtype="pdf",
                           filename="rfq_scan.pdf")
    message.set_boundary("sbam-fixture-boundary")
    write(path, message.as_bytes())


def make_archives(directory: Path) -> None:
    members = {"rfq_summary.md": TEXT_FILES["sample.md"], "bom_extract.csv": TEXT_FILES["sample.csv"]}
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, text in members.items():
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            archive.writestr(info, text)
    write(directory / "sample.zip", zip_buffer.getvalue())

    tar_buffer = io.BytesIO()
    with gzip.GzipFile(filename="", mode="wb", fileobj=tar_buffer, mtime=0) as gz:
        with tarfile.open(fileobj=gz, mode="w") as archive:
            for name, text in members.items():
                data = text.encode("utf-8")
                info = tarfile.TarInfo(name)
                info.size, info.mtime, info.mode = len(data), int(WHEN.timestamp()), 0o644
                archive.addfile(info, io.BytesIO(data))
    write(directory / "sample.tar.gz", tar_buffer.getvalue())


def extract_isa14() -> None:
    """The ISATVON 1.4 sample MS01.1, verbatim from the reference docx (one JSON paragraph)."""
    import docx

    document = docx.Document(REPO / "docs/reference/ISATVON 1.4 for Skill Generation.docx")
    text = next(p.text for p in document.paragraphs if p.text.startswith("{") and "MS01.1" in p.text)
    json.loads(text)  # must be valid JSON
    write(FIX / "isa14/ms01_1_sample.json", text + "\n")


# --- Outlook .msg: an OLE (compound file) container, written by hand ---------------------------------
# A .msg is a Microsoft Compound File Binary (CFB, MS-CFB) holding MAPI properties (MS-OXMSG). The
# writer below is the small part of CFB the fixture needs: version 3 (512-byte sectors), a directory
# tree, a FAT, and a mini stream for streams under 4096 bytes. Nothing is random and every timestamp is
# zero, so the bytes are stable and SHA256SUMS does not change between runs.

FREESECT, ENDOFCHAIN, FATSECT, NOSTREAM = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD, 0xFFFFFFFF
SECTOR, MINI_SECTOR, MINI_CUTOFF = 512, 64, 4096
CFB_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
MSG_CLSID = bytes.fromhex("0b0d0200000000000000c00000000046")  # {00020D0B-0000-0000-C000-000000000046}


class CfbNode:
    """A storage (`data is None`, with children) or a stream (`data` bytes)."""

    def __init__(self, name: str, data: bytes | None = None, children: list["CfbNode"] | None = None):
        self.name, self.data, self.children = name, data, children or []


def _cfb_order(node: CfbNode) -> tuple[int, str]:
    return (len(node.name), node.name.upper())  # MS-CFB: shorter names first, then compare upper case


def _chunks(data: bytes, size: int) -> list[bytes]:
    return [data[i : i + size].ljust(size, b"\0") for i in range(0, len(data), size)]


def _sectors_for(length: int, size: int) -> int:
    return -(-length // size)


def build_cfb(root: CfbNode) -> bytes:
    entries: list[dict] = []  # directory entries, root first

    def add(node: CfbNode, kind: int) -> int:
        entries.append({"node": node, "kind": kind, "left": NOSTREAM, "right": NOSTREAM,
                        "child": NOSTREAM, "start": ENDOFCHAIN, "size": 0})
        return len(entries) - 1

    def link(ids: list[int]) -> int:
        """A balanced binary tree over the siblings (already sorted); returns its root id."""
        if not ids:
            return NOSTREAM
        mid = len(ids) // 2
        entries[ids[mid]]["left"] = link(ids[:mid])
        entries[ids[mid]]["right"] = link(ids[mid + 1 :])
        return ids[mid]

    def walk(node_id: int) -> None:
        node = entries[node_id]["node"]
        kids = sorted(node.children, key=_cfb_order)
        kid_ids = [add(kid, 1 if kid.data is None else 2) for kid in kids]
        entries[node_id]["child"] = link(kid_ids)
        for kid_id in kid_ids:
            walk(kid_id)

    walk(add(root, 5))
    for entry in entries:
        if entry["node"].data is None:
            entry["start"] = 0  # MS-CFB: a storage has no sectors, its start sector is 0

    # Small streams go into the mini stream, big ones get their own sectors.
    mini_stream = b""
    mini_fat: list[int] = []
    big: list[tuple[int, bytes]] = []
    for index, entry in enumerate(entries):
        data = entry["node"].data
        if data is None:
            continue
        entry["size"] = len(data)
        if not data:
            continue
        if len(data) < MINI_CUTOFF:
            first = len(mini_fat)
            count = _sectors_for(len(data), MINI_SECTOR)
            mini_fat += [first + i + 1 for i in range(count - 1)] + [ENDOFCHAIN]
            entry["start"] = first
            mini_stream += data.ljust(count * MINI_SECTOR, b"\0")
        else:
            big.append((index, data))
    entries[0]["size"] = len(mini_stream)

    n_big = sum(_sectors_for(len(d), SECTOR) for _, d in big)
    n_mini_data = _sectors_for(len(mini_stream), SECTOR)
    n_mini_fat = _sectors_for(len(mini_fat) * 4, SECTOR)
    n_dir = _sectors_for(len(entries) * 128, SECTOR)
    n_fat = 1
    while n_fat * (SECTOR // 4) < n_big + n_mini_data + n_mini_fat + n_dir + n_fat:
        n_fat += 1
    assert n_fat <= 109, "a bigger file would need DIFAT sectors"

    fat = [FREESECT] * (n_fat * (SECTOR // 4))
    body: list[bytes] = []

    def chain(first: int, count: int) -> None:
        for i in range(count):
            fat[first + i] = first + i + 1 if i < count - 1 else ENDOFCHAIN

    cursor = 0
    for index, data in big:
        count = _sectors_for(len(data), SECTOR)
        entries[index]["start"] = cursor
        chain(cursor, count)
        body += _chunks(data, SECTOR)
        cursor += count
    if n_mini_data:
        entries[0]["start"] = cursor
        chain(cursor, n_mini_data)
        body += _chunks(mini_stream, SECTOR)
        cursor += n_mini_data
    mini_fat_first = cursor if n_mini_fat else ENDOFCHAIN
    if n_mini_fat:
        chain(cursor, n_mini_fat)
        padded = mini_fat + [FREESECT] * (n_mini_fat * (SECTOR // 4) - len(mini_fat))
        body += _chunks(struct.pack(f"<{len(padded)}I", *padded), SECTOR)
        cursor += n_mini_fat
    dir_first = cursor
    chain(cursor, n_dir)
    directory = b"".join(_pack_entry(e) for e in entries)
    body += _chunks(directory, SECTOR)  # padded with zero entries; fixed up just below
    cursor += n_dir
    fat_first = cursor
    for i in range(n_fat):
        fat[fat_first + i] = FATSECT
    body += _chunks(struct.pack(f"<{len(fat)}I", *fat), SECTOR)

    # Unused directory slots must read as "no stream", not as a stream named "" at sector 0.
    pad = n_dir * 4 - len(entries)
    if pad:
        start = dir_first
        data = bytearray(b"".join(body[start : start + n_dir]))
        for slot in range(len(entries), n_dir * 4):
            offset = slot * 128
            data[offset + 68 : offset + 80] = struct.pack("<III", NOSTREAM, NOSTREAM, NOSTREAM)
        body[start : start + n_dir] = _chunks(bytes(data), SECTOR)

    difat = [fat_first + i for i in range(n_fat)] + [FREESECT] * (109 - n_fat)
    header = (
        CFB_SIGNATURE + bytes(16)
        + struct.pack("<HHHHH", 0x003E, 3, 0xFFFE, 9, 6) + bytes(6)
        + struct.pack("<IIIIIIIII", 0, n_fat, dir_first, 0, MINI_CUTOFF, mini_fat_first, n_mini_fat,
                      ENDOFCHAIN, 0)
        + struct.pack("<109I", *difat)
    )
    assert len(header) == SECTOR
    return header + b"".join(body)


def _pack_entry(entry: dict) -> bytes:
    node: CfbNode = entry["node"]
    raw = node.name.encode("utf-16-le")
    assert len(raw) <= 62, node.name
    clsid = MSG_CLSID if entry["kind"] == 5 else bytes(16)
    return (
        raw.ljust(64, b"\0") + struct.pack("<H", len(raw) + 2)
        + bytes([entry["kind"], 1])  # object type, colour (all black)
        + struct.pack("<III", entry["left"], entry["right"], entry["child"])
        + clsid + bytes(4) + bytes(16)  # CLSID, state bits, created and modified (zero)
        + struct.pack("<IQ", entry["start"], entry["size"])
    )


def _filetime(moment: dt.datetime) -> int:
    return int((moment - dt.datetime(1601, 1, 1, tzinfo=dt.UTC)).total_seconds()) * 10_000_000


def _props(header: bytes, props: list[tuple[int, int, object]]) -> tuple[bytes, list[CfbNode]]:
    """A `__properties_version1.0` stream plus the `__substg1.0_*` streams of the variable-size ones.

    `props` are (property id, type, value); types: 0x0003 long, 0x0040 time, 0x001F text (UTF-16),
    0x0102 binary. A text property's size field counts its two-byte terminator; the stream does not.
    """
    table, streams = b"", []
    for prop_id, kind, value in sorted(props, key=lambda p: p[0]):
        tag = (prop_id << 16) | kind
        if kind == 0x0003:
            table += struct.pack("<IIII", tag, 6, int(value) & 0xFFFFFFFF, 0)  # type: ignore[arg-type]
        elif kind == 0x0040:
            table += struct.pack("<IIQ", tag, 6, _filetime(value))  # type: ignore[arg-type]
        else:
            data = str(value).encode("utf-16-le") if kind == 0x001F else bytes(value)  # type: ignore[arg-type]
            size = len(data) + (2 if kind == 0x001F else 0)
            table += struct.pack("<IIII", tag, 6, size, 0)
            streams.append(CfbNode(f"__substg1.0_{tag:08X}", data))
    return CfbNode("__properties_version1.0", header + table), streams


def build_msg(attachment: bytes) -> bytes:
    """The Outlook message fixture: one sender, one recipient, a body and one PDF attachment."""
    subject = "RFQ ADX-2026-014 with scanned drawing (Outlook)"
    body = "Hello,\r\n\r\nplease find the scanned RFQ attached (synthetic test data).\r\n"
    headers = (
        "From: RFQ Desk <rfq-desk@example.com>\r\nTo: Sourcing <sourcing@example.com>\r\n"
        f"Subject: {subject}\r\nDate: {email.utils.format_datetime(WHEN)}\r\n"
        "Message-ID: <fixture-msg-1@example.com>\r\n"
    )
    message, message_streams = _props(
        bytes(8) + struct.pack("<IIII", 1, 1, 1, 1) + bytes(8),  # next ids, recipient and attach counts
        [
            (0x001A, 0x001F, "IPM.Note"), (0x0037, 0x001F, subject), (0x0039, 0x0040, WHEN),
            (0x0042, 0x001F, "RFQ Desk"), (0x0064, 0x001F, "SMTP"),
            (0x0065, 0x001F, "rfq-desk@example.com"), (0x007D, 0x001F, headers),
            (0x0C1A, 0x001F, "RFQ Desk"), (0x0C1E, 0x001F, "SMTP"),
            (0x0C1F, 0x001F, "rfq-desk@example.com"), (0x0E04, 0x001F, "Sourcing"),
            (0x0E06, 0x0040, WHEN), (0x0E07, 0x0003, 0x11),  # read, has attachment
            (0x1000, 0x001F, body), (0x3FDE, 0x0003, 65001),
        ],
    )
    recipient, recipient_streams = _props(
        bytes(8),
        [
            (0x0C15, 0x0003, 1), (0x3000, 0x0003, 0), (0x3001, 0x001F, "Sourcing"),
            (0x3002, 0x001F, "SMTP"), (0x3003, 0x001F, "sourcing@example.com"),
            (0x39FE, 0x001F, "sourcing@example.com"),
        ],
    )
    attached, attachment_streams = _props(
        bytes(8),
        [
            (0x0E20, 0x0003, len(attachment)), (0x0E21, 0x0003, 0), (0x3001, 0x001F, "rfq_scan.pdf"),
            (0x3701, 0x0102, attachment), (0x3703, 0x001F, ".pdf"), (0x3704, 0x001F, "RFQ_SC~1.PDF"),
            (0x3705, 0x0003, 1), (0x3707, 0x001F, "rfq_scan.pdf"), (0x370B, 0x0003, 0xFFFFFFFF),
            (0x370E, 0x001F, "application/pdf"),
        ],
    )
    root = CfbNode("Root Entry", children=[
        message,
        *message_streams,
        CfbNode("__recip_version1.0_#00000000", children=[recipient, *recipient_streams]),
        CfbNode("__attach_version1.0_#00000000", children=[attached, *attachment_streams]),
        CfbNode("__nameid_version1.0", children=[
            CfbNode("__substg1.0_00020102", b""),  # GUID stream
            CfbNode("__substg1.0_00030102", b""),  # entry stream
            CfbNode("__substg1.0_00040102", b""),  # string stream
        ]),
    ])
    return build_cfb(root)


def make_msg(path: Path, attachment: bytes) -> None:
    write(path, build_msg(attachment))


def copy_masters() -> None:
    name = "BuLLM - Admin Masters (One Sheet) - Consolidated.xlsx"
    shutil.copyfile(REPO / "docs/reference" / name, FIX / "masters" / name)  # unmodified copy


def write_sums() -> None:
    lines = []
    for path in sorted(FIX.rglob("*")):
        relative = path.relative_to(FIX).as_posix()
        if not path.is_file() or relative in NOT_FIXTURES or "__pycache__" in relative:
            continue
        lines.append(f"{hashlib.sha256(path.read_bytes()).hexdigest()}  {relative}")
    write(FIX / "SHA256SUMS", "\n".join(lines) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--sums", action="store_true", help="only rewrite SHA256SUMS")
    parser.add_argument("--msg", action="store_true", help="only rewrite documents/with_attachment.msg")
    args = parser.parse_args()
    if args.msg:
        docs = FIX / "documents"
        make_msg(docs / "with_attachment.msg", (docs / "scanned_full.pdf").read_bytes())
    elif not args.sums:
        docs = FIX / "documents"
        for name, text in TEXT_FILES.items():
            write(docs / name, text)
        page_png, scanned_full = make_pdfs(docs)
        make_images(docs, page_png)
        make_docx(docs / "sample.docx")
        make_xlsx(docs / "multi_sheet.xlsx")
        make_eml(docs / "with_attachment.eml", scanned_full)
        make_msg(docs / "with_attachment.msg", scanned_full)
        make_archives(docs)
        extract_isa14()
        copy_masters()
    write_sums()
    print(f"fixtures written under {FIX.relative_to(REPO)}")


if __name__ == "__main__":
    main()
