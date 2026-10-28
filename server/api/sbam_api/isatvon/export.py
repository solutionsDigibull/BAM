"""GET W/isatvon/export.xlsx (E8.5, C-17): the Workflow's ISATVON breakdown as an Excel workbook.

Export only (C-17: never imported back). The handler answers 302 to a signed URL (300 s); the file
is built here with openpyxl (light work, like the masters workbook export) and stored through
`StorageBackend` under the tenant / project / Workflow `exports` path, keyed by its own sha256.

* The first sheet, `Workflow`, sums it up: name, version, state, the SOW version and its sha256,
  counts (Macro Skills, Sub-Skills, valid, invalid) and the export time.
* Then one sheet per Macro Skill, in order, named by its `macro_category` (by `MSnn` alone when that
  is longer than Excel's 31 characters), with one row per Sub-Skill in order. Deleted Sub-Skills
  are left out; INVALID ones too with `include_invalid=false`.
* The columns are the grid's (`isatvon/grid.py`, E8.6): two header rows, the group (Metadata, I, S,
  A, T, V, O, N, Status) over its columns. A column the metadata gains appears here unchanged.
* A cell holds what the grid shows for its `field_path`: a scalar as it is, an object or array as
  pretty JSON text, null or an empty array as an empty cell. `validation_errors` lists the
  findings, one per line. A cell a finding points at is filled red with a comment (message, code).
* Safety: text starting with `=`, `+`, `-`, `@`, a tab or a carriage return gets a leading
  apostrophe and is never a formula (CSV / Excel injection); text over 32,767 characters (Excel's
  cell limit) is cut with a marker. The model holds secret names only, never values.
"""

import io
import json
import uuid
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any, Final

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sbam_contracts.api import ErrorCode
from sbam_contracts.isatvon14 import FRAMEWORK_VERSION
from sbam_contracts.isatvon_grid import Grid, GridColumn
from sbam_contracts.state_machines import SubSkillState
from sqlalchemy.orm import Session

from sbam_api.audit.repository import record
from sbam_api.core.auth import CurrentUser
from sbam_api.core.conventions import ApiError, not_found
from sbam_api.core.logging import correlation_id_var
from sbam_api.isatvon.grid import render
from sbam_api.isatvon.models import MacroSkill, SubSkill
from sbam_api.isatvon.repository import MacroSkillRepository, SubSkillRepository
from sbam_api.sow.repository import SowVersionRepository
from sbam_api.storage.base import Kind, StorageBackend, StorageScope

XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_CELL: Final = 32_767  # Excel's limit per cell
TRUNCATED: Final = " …[truncated]"
FORMULA_START: Final = ("=", "+", "-", "@", "\t", "\r")
SHEET_CHARS: Final = 31
SUMMARY: Final = "Workflow"
INVALID_FILL: Final = PatternFill("solid", fgColor="FFC7CE")
HEADER_FONT: Final = Font(bold=True)


@dataclass(frozen=True)
class Export:
    storage_key: str
    sha256: str
    size: int


# --- cells ----------------------------------------------------------------------------------------


def safe_text(text: str) -> str:
    """No formula, ever (a leading apostrophe), and within Excel's cell limit."""
    if text.startswith(FORMULA_START):
        text = "'" + text
    if len(text) > MAX_CELL:
        text = text[: MAX_CELL - len(TRUNCATED)] + TRUNCATED
    return text


def cell_value(value: Any) -> Any:
    """What the grid shows: scalars as they are; objects / arrays as pretty JSON; nothing for null
    or an empty array / object (an empty cell, never `[]`)."""
    if value is None or value == [] or value == {}:
        return None
    if isinstance(value, bool | int | float):
        return value
    if isinstance(value, dict | list):
        return safe_text(json.dumps(value, ensure_ascii=False, indent=2))
    return safe_text(str(value))


def value_at(row: SubSkill, path: str) -> Any:
    """The Sub-Skill's value at a grid `field_path` (model paths into `isatvon`, platform fields
    from the row)."""
    if path == "state":
        return row.state
    if path == "content_sha256":
        return row.content_sha256
    if path == "validation_errors":
        return "\n".join(_finding_line(f) for f in row.validation_errors or []) or None
    node: Any = row.isatvon
    for part in path.split("."):
        if not isinstance(node, dict) or part not in node:
            return None
        node = node[part]
    return node


def _finding_line(finding: dict[str, Any]) -> str:
    return f"{finding.get('path', '')}: {finding.get('message', '')} [{finding.get('code', '')}]"


def _column_of(path: str, columns: list[str]) -> str | None:
    """The column a finding's path points into (the longest column path it starts with)."""
    hits = [c for c in columns if path == c or path.startswith(c + ".")]
    return max(hits, key=len) if hits else None


# --- sheets ---------------------------------------------------------------------------------------


def sheet_name(macro_category: str) -> str:
    return macro_category if len(macro_category) <= SHEET_CHARS else macro_category[:4]


def _header(sheet: Worksheet, grid: Grid) -> list[GridColumn]:
    columns: list[GridColumn] = []
    for group in grid.column_groups:
        first = len(columns) + 1
        for column in group.children:
            columns.append(column)
            cell = sheet.cell(row=2, column=len(columns), value=column.header)
            cell.font = HEADER_FONT
            width = min(column.width / 7, 60)
            sheet.column_dimensions[get_column_letter(len(columns))].width = width
        if group.children:
            top = sheet.cell(row=1, column=first, value=group.header)
            top.font = HEADER_FONT
            top.alignment = Alignment(horizontal="center")
            if len(columns) > first:
                last = len(columns)
                sheet.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
    sheet.freeze_panes = "A3"
    return columns


def _rows(sheet: Worksheet, columns: list[GridColumn], rows: Iterable[SubSkill]) -> None:
    paths = [c.field_path for c in columns]
    for r, row in enumerate(rows, start=3):
        for c, path in enumerate(paths, start=1):
            sheet.cell(row=r, column=c, value=cell_value(value_at(row, path)))
        for finding in row.validation_errors or []:
            target = _column_of(str(finding.get("path", "")), paths)
            if target is None:
                continue
            cell = sheet.cell(row=r, column=paths.index(target) + 1)
            cell.fill = INVALID_FILL
            note = f"{finding.get('message', '')} [{finding.get('code', '')}]"
            if cell.comment is not None:
                note = f"{cell.comment.text}\n{note}"
            cell.comment = Comment(safe_text(note), "SBAM")


def _summary(sheet: Worksheet, facts: list[tuple[str, Any]]) -> None:
    for r, (label, value) in enumerate(facts, start=1):
        sheet.cell(row=r, column=1, value=label).font = HEADER_FONT
        sheet.cell(row=r, column=2, value=cell_value(value))
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 70


def build(
    workflow: Any,
    sow_version: Any,
    macros: list[tuple[MacroSkill, list[SubSkill]]],
    grid: Grid,
    exported_at: datetime,
) -> bytes:
    """The workbook bytes (pure: no database, no clock)."""
    book = openpyxl.Workbook()
    summary = book.create_sheet(SUMMARY, 0)
    del book[book.sheetnames[1]]  # the default empty sheet
    rows = [r for _, subs in macros for r in subs]
    _summary(
        summary,
        [
            ("Workflow", workflow.name),
            ("Version", workflow.version),
            ("State", workflow.state),
            ("SOW version", sow_version.version if sow_version is not None else None),
            ("SOW sha256", sow_version.sha256 if sow_version is not None else None),
            ("Macro Skills", len(macros)),
            ("Sub-Skills", len(rows)),
            ("Valid", sum(r.state == SubSkillState.VALID.value for r in rows)),
            ("Invalid", sum(r.state == SubSkillState.INVALID.value for r in rows)),
            ("ISATVON", FRAMEWORK_VERSION),
            ("Exported at", exported_at.isoformat(timespec="seconds")),
        ],
    )
    for macro, subs in macros:
        sheet = book.create_sheet(sheet_name(macro.macro_category))
        _rows(sheet, _header(sheet, grid), subs)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# --- the request ----------------------------------------------------------------------------------


def export_workbook(
    session: Session,
    storage: StorageBackend,
    user: CurrentUser,
    project_id: uuid.UUID,
    workflow_id: uuid.UUID,
    *,
    include_invalid: bool = True,
) -> Export:
    """Build and store the workbook of Workflow `workflow_id`; audited as ISATVON_EXPORTED."""
    from sbam_api.workflows.repository import WorkflowRepository  # the SOW package imports this

    ctx = user.context(project_id)
    workflow = WorkflowRepository(session, ctx).get(workflow_id)
    if workflow is None:
        raise not_found()
    macro_rows = MacroSkillRepository(session, ctx).of_workflow(workflow.id)
    if not macro_rows:
        raise ApiError(
            409,
            ErrorCode.CONFLICT,
            "Nothing to export",
            "The Workflow has no Macro Skills yet; export once the breakdown is extracted.",
        )
    subs = SubSkillRepository(session, ctx)
    macros = [
        (
            m,
            [
                r
                for r in subs.of_macro(m.id)
                if include_invalid or r.state != SubSkillState.INVALID.value
            ],
        )
        for m in macro_rows
    ]
    sow_version = SowVersionRepository(session, ctx).get(workflow.sow_version_id)
    rendered = render(FRAMEWORK_VERSION)
    if rendered is None:  # never: the served framework version always renders
        raise not_found()
    data = build(workflow, sow_version, macros, rendered.payload.grid, datetime.now(UTC))
    stored = storage.put(
        data, StorageScope(ctx.tenant_id, project_id, workflow.id), Kind.EXPORTS, XLSX
    )
    record(
        session,
        tenant_id=ctx.tenant_id,
        actor_id=str(user.user_id),
        action="ISATVON_EXPORTED",
        entity_type="workflow",
        entity_id=str(workflow.id),
        after={
            "sha256": stored.sha256,
            "macro_skills": len(macros),
            "sub_skills": sum(len(r) for _, r in macros),
            "include_invalid": include_invalid,
        },
        correlation_id=correlation_id_var.get() or "isatvon-export",
    )
    session.commit()
    return Export(stored.storage_key, stored.sha256, stored.size)
