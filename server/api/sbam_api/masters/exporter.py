"""Masters workbook export and empty template (TD §8.3, E3.4): `GET /admin/masters/export|template`.

Both start from `template.xlsx`, the frozen BAM 0.2 layout (README, header rows 1-3 and styling of
every tab, the 43 frozen Filetype rows), built by `scripts/build_masters_template.py`. The export
writes the caller's tenant entries (plus, with `project_id`, that project's own) from row 4 on,
in the column order the importer reads, so import -> export -> import changes nothing:
  * row labels and payload values only; SBAM metadata (stable id, status) stays in the Misc text;
  * every value is a string cell, never a formula, so `=HYPERLINK(...)` round-trips as text;
  * the Filetype tab is the platform's frozen one, whatever the tenant stored;
  * the Tooling tab is kept only when there are Tooling entries; the template never has it.
Restricted entries are exported too: only the tenant's admins can call this.

The same file answers `GET /admin/master-types` (TD §8.3): the column headers and row-3
descriptions the masters screen shows, so the workbook stays the one source of the layout.
"""

import io
import json
import re
import uuid
from datetime import UTC, datetime
from functools import cache
from importlib.resources import files
from typing import Any, Final

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from sbam_contracts.api import ErrorCode, FieldError
from sbam_contracts.masters import (
    ADMIN_EDITABLE,
    MAX_CELL_CHARS,
    README_MAX_COLUMNS,
    MasterType,
    ReadmeTab,
    cell_text,
)
from sbam_contracts.masters_api import MasterColumn, MasterTypeInfo, MasterTypes
from sbam_contracts.masters_workbook import (
    COLUMNS,
    HEADER_ROWS,
    PAYLOAD_MODELS,
    README_TAB,
    ROW_LABEL_PREFIX,
    TABS,
    TOOLING_TAB,
)
from sqlalchemy.orm import Session

from sbam_api.audit.repository import record
from sbam_api.core.auth import CurrentUser
from sbam_api.core.conventions import ApiError, not_found
from sbam_api.identity.models import Tenant
from sbam_api.masters.service import _correlation, _repo
from sbam_api.projects.repository import ProjectRepository

XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TENANT_TYPES: Final = [t for t in TABS if t != MasterType.FILETYPE]


def template() -> bytes:
    book = _layout()
    del book[TOOLING_TAB]
    return _save(book)


def export(session: Session, user: CurrentUser, project_id: uuid.UUID | None) -> tuple[str, bytes]:
    """(file name, workbook bytes) of the caller's masters, audited as MASTERS_EXPORTED."""
    if (
        project_id is not None
        and ProjectRepository(session, user.context()).get(project_id) is None
    ):
        raise not_found()
    repo = _repo(session, user)
    book = _layout()
    counts: dict[str, int] = {}
    for master_type in TENANT_TYPES:
        entries = repo.of_type(master_type.value, project_id)
        versions = repo.current_versions(entries)
        rows = sorted(
            ((e.row_label, versions[e.id].payload) for e in entries if e.id in versions),
            key=lambda row: _label_order(row[0]),
        )
        tab = TABS[master_type]
        if master_type == MasterType.TOOLING and not rows:
            del book[tab]
            continue
        _write_rows(book[tab], master_type, rows)
        counts[tab] = len(rows)
    tenant = session.get(Tenant, user.tenant_id)
    slug = re.sub(r"[^A-Za-z0-9]+", "-", tenant.name if tenant else "").strip("-") or "tenant"
    name = f"BAM_masters_{slug}_{datetime.now(UTC):%Y%m%d}.xlsx"
    record(
        session,
        tenant_id=user.tenant_id,
        actor_id=str(user.user_id),
        action="MASTERS_EXPORTED",
        entity_type="masters_export",
        entity_id=str(project_id) if project_id else None,
        after={
            "filename": name,
            "project_id": str(project_id) if project_id else None,
            **counts,
        },
        correlation_id=_correlation(),
    )
    data = _save(book)
    session.commit()
    return name, data


@cache
def master_types() -> MasterTypes:
    """Every master's frozen columns with their row-3 description, and the README tab."""
    book = _layout()
    types = []
    for master_type, tab in TABS.items():
        row3 = next(book[tab].iter_rows(min_row=HEADER_ROWS, max_row=HEADER_ROWS, values_only=True))
        descriptions = [cell_text(v).strip() for v in row3[1:]]  # column A is the row label
        columns = [
            MasterColumn(header=header, field=field, description=descriptions[i])
            for i, (header, field) in enumerate(COLUMNS[master_type])
        ]
        types.append(
            MasterTypeInfo(
                master_type=master_type,
                tab=tab,
                editable=master_type in ADMIN_EDITABLE,
                row_label_prefix=ROW_LABEL_PREFIX[master_type],
                columns=columns,
                payload_schema=PAYLOAD_MODELS[master_type].model_json_schema(by_alias=True),
            )
        )
    rows = [
        [None if v is None else cell_text(v) for v in row[:README_MAX_COLUMNS]]
        for row in book[README_TAB].iter_rows(values_only=True)
    ]
    while rows and not any(rows[-1]):
        rows.pop()
    readme = ReadmeTab(title=rows[0][0] or "", rows=rows[1:])
    return MasterTypes(types=types, readme=readme)


def _layout() -> openpyxl.Workbook:
    data = files("sbam_api.masters").joinpath("template.xlsx").read_bytes()
    return openpyxl.load_workbook(io.BytesIO(data))


def _save(book: openpyxl.Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _label_order(label: str) -> tuple[int, str]:
    """`Rule #2` before `Rule #10`; Tooling ids (no number) by name."""
    number = re.search(r"#\s*(\d+)", label)
    return (int(number[1]) if number else 0, label)


def _cell_text(field: str, value: Any) -> str:
    if value is None:
        return ""
    if field == "rag_reference":  # stored with the cell's own keys (`Links / Path`, ...)
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _write_rows(
    sheet: Worksheet, master_type: MasterType, rows: list[tuple[str, dict[str, Any]]]
) -> None:
    for sheet_row, (label, payload) in enumerate(rows, start=HEADER_ROWS + 1):
        texts = [label] + [_cell_text(f, payload.get(f)) for _, f in COLUMNS[master_type]]
        for column, text in enumerate(texts, start=1):
            if len(text) > MAX_CELL_CHARS:  # e.g. a RAG Reference whose JSON outgrows one cell
                path = f"{sheet.title}!row{sheet_row}:col{column}"
                message = f"{label} has a value longer than {MAX_CELL_CHARS} characters"
                raise ApiError(
                    422,
                    ErrorCode.VALIDATION_ERROR,
                    "A value is too long for an Excel cell",
                    f"Shorten {label} on the {sheet.title} tab, then export again.",
                    errors=[FieldError(path=path, message=message)],
                )
            if text:
                cell = sheet.cell(sheet_row, column, text)
                cell.data_type = "s"  # openpyxl reads a leading '=' as a formula otherwise
