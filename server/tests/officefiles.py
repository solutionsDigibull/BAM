"""Small .docx and .xlsx files for tests, built with python-docx and openpyxl.

Shared by the unit and the integration tests (nothing is committed to the fixtures folder).
"""

import io
from typing import Any

from docx import Document
from openpyxl import Workbook


def docx_bytes(build: Any) -> bytes:
    doc = Document()
    build(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def sample_docx() -> bytes:
    def build(doc: Any) -> None:
        doc.add_heading("Quote terms", 0)
        doc.add_heading("Scope", 1)
        doc.add_paragraph("Supply of 2,500 boards.")
        doc.add_paragraph("IPC Class 3", style="List Bullet")
        doc.add_paragraph("AS9100D", style="List Bullet")
        doc.add_heading("Prices", 2)
        table = doc.add_table(rows=2, cols=2)
        for cell, text in zip(table._cells, ["MPN", "Qty", "STM32", "12"], strict=True):
            cell.text = text
        doc.add_paragraph("Valid for 30 days.")

    return docx_bytes(build)


def sheet_of(book: Workbook) -> Any:
    """The active sheet (typed Optional by openpyxl; a new workbook always has one)."""
    return book.active


def xlsx_bytes(build: Any) -> bytes:
    book = Workbook()
    build(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
