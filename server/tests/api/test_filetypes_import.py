"""E3.13 super-admin Filetype import through the real app and auth chain, on SQLite, with a fake
clamd. Platform rows change only through POST /admin/filetypes/import (ADR 0020)."""

import hashlib
import uuid
from pathlib import Path
from typing import Any

import pytest
import test_projects
from clamd_fake import FakeClamd
from openpyxl.workbook.workbook import Workbook
from sbam_api.filetypes.config import CONFIG_FILE
from sbam_api.filetypes.models import FiletypeRowVersion
from sqlalchemy import delete, select, update
from sqlalchemy.exc import DatabaseError
from test_auth import problem
from test_masters_api import engine_with_masters
from test_masters_import import XLSX, ImportHarness, clamd, column, tab, workbook

__all__ = ["clamd"]  # the fixture, used by name

ROW_7 = 10  # sheet row of Filetype #7


class FiletypeHarness(ImportHarness):
    def root(self) -> str:
        return self.realm.token(
            uuid.uuid4(), tenant_id=None, realm_access={"roles": ["super-admin"]}
        )

    def platform(self, data: bytes, mode: str = "upsert", **form: str) -> Any:
        return self.call(
            "POST",
            "/admin/filetypes/import",
            self.root(),
            files={"file": ("filetypes.xlsx", data, XLSX)},
            data={"mode": mode, **form},
        )

    def versions(self) -> list[FiletypeRowVersion]:
        with self.sessions() as s:
            return list(s.scalars(select(FiletypeRowVersion)))

    def tier_7(self) -> int:
        res = self.call("GET", "/admin/filetypes", self.root())
        assert res.status_code == 200, res.text
        tier: int = res.json()[6]["tier"]
        return tier


@pytest.fixture
def h(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, clamd: FakeClamd) -> FiletypeHarness:
    monkeypatch.setattr(test_projects, "sqlite_engine", engine_with_masters)
    h = FiletypeHarness(tmp_path)
    state = h.client.app.state  # type: ignore[attr-defined]
    state.settings = state.settings.model_copy(
        update={"clamav_host": clamd.host, "clamav_port": clamd.port, "clamav_timeout": 5.0}
    )
    return h


def retier(book: Workbook) -> None:
    cell = book["Filetype"].cell(ROW_7, column(book, "Filetype", "Tier"))
    cell.value = 3 if cell.value != 3 else 2


def test_e3_13_t01_super_admin_changes_filetype_7_tier(h: FiletypeHarness) -> None:
    before = h.tier_7()
    res = h.platform(workbook(retier), version_label="BAM 0.3")
    assert res.status_code == 200, res.text
    report = res.json()
    assert report["applied"] is True
    assert (tab(report, "Filetype")["updated"], tab(report, "Filetype")["unchanged"]) == (1, 42)

    [version] = h.versions()
    assert (version.row_label, version.version) == ("Filetype #7", 2)
    assert (version.version_label, version.workbook_sha256) == ("BAM 0.3", report["sha256"])

    [audit] = h.actions("FILETYPE_VERSION_IMPORTED")
    assert audit.tenant_id is None and audit.entity_id == report["sha256"]
    assert audit.before is not None and audit.after is not None
    assert audit.after["changed"] == ["Filetype #7"]
    assert audit.before["Filetype #7"]["tier"] == before
    assert audit.after["rows"]["Filetype #7"]["tier"] == h.tier_7() != before


def test_e3_13_t01_validate_only_writes_nothing(h: FiletypeHarness) -> None:
    res = h.platform(workbook(retier), mode="validate-only")
    assert res.status_code == 200 and res.json()["applied"] is False
    assert tab(res.json(), "Filetype")["updated"] == 1
    assert h.versions() == [] and h.actions("FILETYPE_VERSION_IMPORTED") == []


def test_e3_13_t02_tenant_admin_cannot_change_filetype(h: FiletypeHarness) -> None:
    problem(h.upload(workbook(retier)), 403, "FILETYPE_FROZEN")
    assert h.counts() == (0, 0) and h.versions() == []
    res = h.call(
        "POST",
        "/admin/filetypes/import",
        h.admin_a(),
        files={"file": ("filetypes.xlsx", workbook(retier), XLSX)},
        data={"mode": "upsert"},
    )
    problem(res, 403, "WRONG_PORTAL")
    assert h.versions() == []


def test_e3_13_t03_unchanged_and_repeated_imports_write_nothing(h: FiletypeHarness) -> None:
    res = h.platform(workbook())
    assert res.status_code == 200 and tab(res.json(), "Filetype")["unchanged"] == 43
    assert h.versions() == [] and h.actions("FILETYPE_VERSION_IMPORTED") == []

    changed = workbook(retier)
    assert h.platform(changed).status_code == 200
    again = h.platform(changed)
    assert again.status_code == 200 and tab(again.json(), "Filetype")["updated"] == 0
    assert len(h.versions()) == 1 and len(h.actions("FILETYPE_VERSION_IMPORTED")) == 1


def test_e3_13_t03_second_change_is_version_3(h: FiletypeHarness) -> None:
    h.platform(workbook(retier))
    assert h.platform(workbook()).status_code == 200  # back to the shipped values
    assert sorted(v.version for v in h.versions()) == [2, 3]


def test_e3_13_t04_44th_row_is_422(h: FiletypeHarness) -> None:
    def extra(book: Workbook) -> None:
        sheet = book["Filetype"]
        last = sheet.max_row
        sheet.append([c.value for c in sheet[last]])
        sheet.cell(last + 1, 1).value = "Filetype #44"

    problem(h.platform(workbook(extra)), 422, "STRUCTURE_MISMATCH")
    assert h.versions() == []


def test_e3_13_t05_loader_config_is_untouched(h: FiletypeHarness) -> None:
    config = CONFIG_FILE.read_bytes()
    assert h.platform(workbook(retier)).status_code == 200
    assert CONFIG_FILE.read_bytes() == config
    [version] = h.versions()
    assert set(version.payload) == {"file_format", "tier", "purpose", "rag_ingestion", "example"}


def test_e3_13_tenant_import_compares_to_the_platform_rows(h: FiletypeHarness) -> None:
    assert h.platform(workbook(retier)).status_code == 200
    problem(h.upload(workbook()), 403, "FILETYPE_FROZEN")  # the old values
    assert h.counts() == (0, 0)
    res = h.upload(workbook(retier))  # the new platform values
    assert res.status_code == 200, res.text
    assert tab(res.json(), "Filetype")["created"] == 43


def test_e3_13_versions_are_append_only(h: FiletypeHarness) -> None:
    h.platform(workbook(retier))
    for statement in (
        update(FiletypeRowVersion).values(version=9),
        delete(FiletypeRowVersion),
    ):
        with h.sessions() as s, pytest.raises(DatabaseError):
            s.execute(statement)
    assert len(h.versions()) == 1


def test_e3_13_format_change_on_an_enabled_row_is_applied_with_a_warning(
    h: FiletypeHarness,
) -> None:
    def reformat(book: Workbook) -> None:  # Filetype #1 (.cpp) is enabled for loading
        book["Filetype"].cell(4, column(book, "Filetype", "File Format")).value = ".cc"

    res = h.platform(workbook(reformat))
    assert res.status_code == 200, res.text
    [audit] = h.actions("FILETYPE_VERSION_IMPORTED")
    assert audit.after is not None
    assert audit.after["warnings"] == [
        "Filetype #1 is enabled for loading and its file format changed"
    ]
    res = h.call("GET", "/admin/filetypes", h.root())
    assert (res.json()[0]["file_format"], res.json()[0]["loader"]) == (".cc", "source-code")


def test_e3_13_infected_workbook_is_quarantined_under_the_platform_scope(
    h: FiletypeHarness, clamd: FakeClamd
) -> None:
    clamd.behaviour = "infected"
    data = workbook(retier)
    problem(h.platform(data), 422, "MALWARE")
    zero = "00000000-0000-0000-0000-000000000000"
    key = f"{zero}/{zero}/quarantine/{hashlib.sha256(data).hexdigest()}"
    root: Path = h.client.app.state.settings.storage_root  # type: ignore[attr-defined]
    assert (root / key).read_bytes() == data
    [audit] = h.actions("MASTERS_IMPORT_REJECTED")
    assert audit.tenant_id is None and audit.after is not None
    assert audit.after["quarantine_key"] == key
    assert h.versions() == []
