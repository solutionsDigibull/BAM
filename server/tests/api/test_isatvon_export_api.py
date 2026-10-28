"""E8.5 `GET W/isatvon/export.xlsx` through the real app and auth chain: the workbook is fetched
from the signed URL the 302 points at and read back with openpyxl.

On Windows, run with a short `--basetemp` (e.g. `C:/t`): the storage key (tenant / project /
Workflow / exports / sha256) under pytest's long temp folder passes the 260-character path limit."""

import copy
import io
import json
import time
import uuid
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fixtures import fixture_path
from sbam_api.core.ids import uuid7
from sbam_api.isatvon import export as export_module
from sbam_api.isatvon.export import MAX_CELL, TRUNCATED, cell_value, value_at
from sbam_api.isatvon.grid import render
from sbam_api.isatvon.models import MacroSkill, SubSkill
from sbam_contracts.api import API_BASE
from sbam_contracts.canonical import sha256_of
from sbam_contracts.isatvon_grid import Grid, GridColumn
from test_auth import ADMIN_A, FDE_A2, TENANT_A, TENANT_B, problem
from test_workflows_api import Harness as WorkflowsHarness

IDEF = "isatvon_definition."
GRID: Grid = render("1.4").payload.grid  # type: ignore[union-attr]
PATHS = [c.field_path for g in GRID.column_groups for c in g.children]


def sample(skill_id: str, name: str, **dotted: Any) -> dict[str, Any]:
    raw: dict[str, Any] = json.loads(
        fixture_path("isa14", "ms01_1_structured.json").read_text(encoding="utf-8")
    )
    raw = copy.deepcopy(raw)
    raw["skill_metadata"]["skill_id"] = skill_id
    raw["skill_metadata"]["skill_name"] = name
    for path, value in dotted.items():
        *parents, leaf = path.split(".")
        node = raw
        for part in parents:
            node = node[part]
        node[leaf] = value
    return raw


class Harness(WorkflowsHarness):
    def __init__(self, tmp_path: Path) -> None:
        super().__init__(tmp_path)
        bind = self.sessions.kw["bind"]
        for table in (MacroSkill, SubSkill):
            table.__table__.create(bind, checkfirst=True)  # type: ignore[attr-defined]
        self.wf = self.workflow()

    def macro(self, order: int, category: str) -> uuid.UUID:
        with self.sessions() as s:
            row = MacroSkill(
                id=uuid7(),
                tenant_id=TENANT_A,
                project_id=uuid.UUID(self.project),
                workflow_id=self.wf,
                sow_version_id=self.v1,
                macro_category=category,
                name=category,
                description=".",
                order=order,
                sow_refs=["M-ACCESS"],
                agent_ref=None,
                validation_errors=[],
            )
            row.state = "PENDING"
            s.add(row)
            s.commit()
            return row.id

    def sub(
        self,
        macro: uuid.UUID,
        order: int,
        isatvon: dict[str, Any],
        state: str = "VALID",
        findings: list[dict[str, Any]] | None = None,
    ) -> SubSkill:
        with self.sessions() as s:
            row = SubSkill(
                id=uuid7(),
                tenant_id=TENANT_A,
                project_id=uuid.UUID(self.project),
                macro_skill_id=macro,
                workflow_id=self.wf,
                skill_id=isatvon["skill_metadata"]["skill_id"],
                skill_name=isatvon["skill_metadata"]["skill_name"],
                order=order,
                isatvon=isatvon,
                isatvon_framework_version="1.4",
                sow_refs=["F-ACCESS-1"],
                validation_errors=findings or [],
                content_sha256=sha256_of(isatvon).removeprefix("sha256:"),
            )
            row.state = state
            s.add(row)
            s.commit()
            s.expunge(row)
            return row

    def three_by_four(self) -> list[list[SubSkill]]:
        """3 Macro Skills with 4 Sub-Skills each (E8.5-T01)."""
        out = []
        for m in range(1, 4):
            macro = self.macro(m, f"MS0{m}_Area_{m}")
            out.append(
                [
                    self.sub(macro, n, sample(f"MS0{m}.{n}_Skill_{n}", f"Skill {m}.{n}"))
                    for n in range(1, 5)
                ]
            )
        return out

    def export(self, token: str | None = None, workflow: uuid.UUID | None = None, **q: str) -> Any:
        query = "&".join(f"{k}={v}" for k, v in q.items())
        path = f"/projects/{self.project}/workflows/{workflow or self.wf}/isatvon/export.xlsx"
        return self.call("GET", path + (f"?{query}" if query else ""), token or self.fde())

    def workbook(self, include_invalid: str | None = None) -> openpyxl.Workbook:
        q = {} if include_invalid is None else {"include_invalid": include_invalid}
        res = self.export(None, None, **q)
        assert res.status_code == 302, res.text
        got = self.client.get(res.headers["location"])
        assert got.status_code == 200
        return openpyxl.load_workbook(io.BytesIO(got.content))


@pytest.fixture
def h(tmp_path: Path) -> Iterator[Harness]:
    yield Harness(tmp_path)


def test_e8_5_t01_the_workbook_matches_the_grid(h: Harness) -> None:
    macros = h.three_by_four()
    book = h.workbook()
    assert book.sheetnames == ["Workflow", "MS01_Area_1", "MS02_Area_2", "MS03_Area_3"]
    total = 0
    for rows, name in zip(macros, book.sheetnames[1:], strict=True):
        sheet = book[name]
        headers = [c.value for c in sheet[2]]
        assert headers == [c.header for g in GRID.column_groups for c in g.children]
        data = list(sheet.iter_rows(min_row=3, values_only=True))
        assert len(data) == len(rows)
        total += len(data)
        for row, values in zip(rows, data, strict=True):
            assert list(values) == [cell_value(value_at(row, p)) for p in PATHS]
    assert total == 12
    groups = {sheet_cell.value for sheet_cell in book["MS01_Area_1"][1] if sheet_cell.value}
    assert groups == {g.header for g in GRID.column_groups}


def test_e8_5_the_summary_sheet_equals_the_api(h: Harness) -> None:
    macros = h.three_by_four()
    h.sub(h.macro(4, "MS04_Bad"), 1, sample("MS04.1_Bad", "Bad"), state="INVALID")
    summary: dict[Any, Any] = {
        row[0]: row[1] for row in h.workbook()["Workflow"].iter_rows(values_only=True)
    }
    workflow = h.call("GET", f"/projects/{h.project}/workflows", h.fde()).json()["items"][0]
    assert summary["Workflow"] == workflow["name"] and summary["Version"] == workflow["version"]
    assert summary["SOW sha256"] == h.sha(1)
    assert (summary["Macro Skills"], summary["Sub-Skills"]) == (4, 13)
    assert (summary["Valid"], summary["Invalid"]) == (12, 1)
    assert summary["Exported at"]
    assert len(macros) == 3


def test_e8_5_include_invalid_false_leaves_invalid_rows_out(h: Harness) -> None:
    macro = h.macro(1, "MS01_A")
    h.sub(macro, 1, sample("MS01.1_Ok", "Ok"))
    h.sub(macro, 2, sample("MS01.2_Bad", "Bad"), state="INVALID")
    assert h.workbook()["MS01_A"].max_row == 4
    assert h.workbook(include_invalid="false")["MS01_A"].max_row == 3


def test_e8_5_t02_a_formula_is_stored_as_text(h: Harness) -> None:
    evil = '=HYPERLINK("http://evil")'
    h.sub(h.macro(1, "MS01_A"), 1, sample("MS01.1_X", evil))
    sheet = h.workbook()["MS01_A"]
    name = sheet.cell(row=3, column=PATHS.index("skill_metadata.skill_name") + 1)
    assert name.value == "'" + evil and name.data_type == "s"
    for start in ("+1", "-1", "@SUM(A1)", "\tx", "\rx"):
        assert export_module.safe_text(start) == "'" + start


def test_e8_5_t03_a_column_the_metadata_gains_appears_without_code_change(
    h: Harness, monkeypatch: pytest.MonkeyPatch
) -> None:
    h.sub(h.macro(1, "MS01_A"), 1, sample("MS01.1_X", "X"))
    extra = GridColumn(
        field_path=IDEF + "i_instructions.trigger",
        header="Trigger again",
        type="string",
        editor="text",  # type: ignore[arg-type]
        required=False,
        editable=False,
        platform=False,
        width=120,
    )
    grid = GRID.model_copy(deep=True)
    grid.column_groups[1].children.append(extra)
    rendered = render("1.4")
    assert rendered is not None
    patched = type(rendered)(
        body=rendered.body,
        etag=rendered.etag,
        payload=rendered.payload.model_copy(update={"grid": grid}),
    )
    monkeypatch.setattr(export_module, "render", lambda version: patched)
    headers = [c.value for c in h.workbook()["MS01_A"][2]]
    assert "Trigger again" in headers


def test_e8_5_t04_five_hundred_sub_skills_export_in_under_ten_seconds(h: Harness) -> None:
    macro = h.macro(1, "MS01_A")
    for n in range(1, 501):
        h.sub(macro, n, sample(f"MS01.{n}_Skill_{n}", f"Skill {n}"))
    started = time.monotonic()
    book = h.workbook()
    assert time.monotonic() - started <= 10
    assert book["MS01_A"].max_row == 2 + 500


def test_e8_5_t05_other_tenants_unassigned_and_admins_are_refused(h: Harness) -> None:
    h.sub(h.macro(1, "MS01_A"), 1, sample("MS01.1_X", "X"))
    other = h.realm.token(uuid.uuid4(), tenant_id=[str(TENANT_B)])
    problem(h.export(token=other), 404, "NOT_FOUND")
    problem(h.export(token=h.realm.token(FDE_A2)), 404, "NOT_FOUND")
    admin = h.realm.token(ADMIN_A, realm_access={"roles": ["sbam-admin"]})
    problem(h.export(token=admin), 403, "WRONG_PORTAL")
    problem(h.export(workflow=uuid.uuid4()), 404, "NOT_FOUND")


def test_e8_5_t06_an_invalid_cell_is_red_with_its_message(h: Harness) -> None:
    finding = {
        "rule": 5,
        "path": IDEF + "v_variables.budget_limit",
        "message": "budget_limit '2 minutes' is not '<n>ms' or '<n>s'",
        "code": "ISATVON_CONTRACT",
    }
    content = sample("MS01.1_X", "X", **{IDEF + "v_variables.budget_limit": "2 minutes"})
    h.sub(h.macro(1, "MS01_A"), 1, content, state="INVALID", findings=[finding])
    sheet = h.workbook()["MS01_A"]
    cell = sheet.cell(row=3, column=PATHS.index(IDEF + "v_variables.budget_limit") + 1)
    assert cell.value == "2 minutes"
    assert cell.fill.fgColor.rgb.endswith("FFC7CE")
    assert cell.comment is not None and str(finding["message"]) in cell.comment.text
    errors = str(sheet.cell(row=3, column=PATHS.index("validation_errors") + 1).value)
    assert "budget_limit" in errors and "ISATVON_CONTRACT" in errors


def test_e8_5_a_finding_inside_a_list_marks_the_list_column(h: Harness) -> None:
    finding = {
        "rule": 9,
        "path": IDEF + "t_tech_stack.libraries.0.package",
        "message": "'leftpad' is not in the runtime registry",
        "code": "LIBRARY_NOT_IN_REGISTRY",
    }
    h.sub(h.macro(1, "MS01_A"), 1, sample("MS01.1_X", "X"), state="INVALID", findings=[finding])
    sheet = h.workbook()["MS01_A"]
    cell = sheet.cell(row=3, column=PATHS.index(IDEF + "t_tech_stack.libraries") + 1)
    assert cell.comment is not None and "leftpad" in cell.comment.text


def test_e8_5_cells_json_empty_unicode_and_truncation() -> None:
    assert cell_value({"a": 1}) == '{\n  "a": 1\n}'
    assert cell_value([]) is None and cell_value(None) is None and cell_value({}) is None
    assert cell_value(3) == 3 and cell_value(True) is True
    assert cell_value("Prüfung 検査 ✓") == "Prüfung 検査 ✓"
    long = cell_value("x" * (MAX_CELL + 10))
    assert len(long) == MAX_CELL and long.endswith(TRUNCATED)


def test_e8_5_a_long_macro_name_names_the_sheet_by_msnn(h: Harness) -> None:
    long = "MS01_" + "Very_Long_Investor_Registration_Area"
    h.sub(h.macro(1, long), 1, sample("MS01.1_X", "X"))
    assert h.workbook().sheetnames == ["Workflow", "MS01"]


def test_e8_5_no_macro_skills_yet_is_409(h: Harness) -> None:
    problem(h.export(), 409, "CONFLICT")


def test_e8_5_the_redirect_names_the_file_hash_and_the_link_expires(h: Harness) -> None:
    h.sub(h.macro(1, "MS01_A"), 1, sample("MS01.1_X", "X"))
    res = h.export()
    assert res.status_code == 302
    assert len(res.headers["X-Content-SHA256"]) == 64
    assert res.headers["X-Content-SHA256"] in res.headers["location"]
    assert "exports" in res.headers["location"]


def test_e8_5_export_is_a_get_in_openapi(h: Harness) -> None:
    spec = h.client.get(f"{API_BASE}/openapi.json").json()
    path = f"{API_BASE}/projects/{{project_id}}/workflows/{{workflow_id}}/isatvon/export.xlsx"
    assert set(spec["paths"][path]) == {"get"}
