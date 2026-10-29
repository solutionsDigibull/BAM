"""E3.3 masters workbook import through the real app and auth chain, on SQLite, with a fake clamd.

The consolidated BAM 0.2 workbook from the fixture pack is the valid upload; the bad ones are made
from it with openpyxl (or by rewriting its zip) in the test.
"""

import hashlib
import io
import json
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any
from unittest import mock

import openpyxl
import pytest
import test_projects
from clamd_fake import FakeClamd
from fixtures import eicar_bytes, fixture_path
from openpyxl.workbook.workbook import Workbook
from sbam_api.audit.models import AuditLog
from sbam_api.db.repository import StaleRowVersion
from sbam_api.masters.models import MasterEntry, MasterEntryVersion
from sqlalchemy import func, select
from test_auth import TENANT_A, problem
from test_masters_api import MastersHarness, engine_with_masters

WORKBOOK = fixture_path("masters", "BuLLM - Admin Masters (One Sheet) - Consolidated.xlsx")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ImportHarness(MastersHarness):
    def upload(
        self,
        data: bytes,
        mode: str = "upsert",
        filename: str = "masters.xlsx",
        token: str | None = None,
        **form: str,
    ) -> Any:
        return self.call(
            "POST",
            "/admin/masters/import",
            token or self.admin_a(),
            files={"file": (filename, data, XLSX)},
            data={"mode": mode, **form},
        )

    def counts(self) -> tuple[int, int]:
        with self.sessions() as s:
            entries = s.scalar(select(func.count()).select_from(MasterEntry))
            versions = s.scalar(select(func.count()).select_from(MasterEntryVersion))
        return entries or 0, versions or 0

    def actions(self, action: str) -> list[AuditLog]:
        with self.sessions() as s:
            return list(s.scalars(select(AuditLog).where(AuditLog.action == action)))


@pytest.fixture
def clamd() -> Iterator[FakeClamd]:
    with FakeClamd() as fake:
        yield fake


def harness(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, clamd: FakeClamd) -> ImportHarness:
    monkeypatch.setattr(test_projects, "sqlite_engine", engine_with_masters)
    h = ImportHarness(tmp_path)
    state = h.client.app.state  # type: ignore[attr-defined]
    state.settings = state.settings.model_copy(
        update={"clamav_host": clamd.host, "clamav_port": clamd.port, "clamav_timeout": 5.0}
    )
    return h


@pytest.fixture
def h(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, clamd: FakeClamd) -> ImportHarness:
    return harness(tmp_path, monkeypatch, clamd)


def workbook(change: Callable[[Workbook], None] | None = None) -> bytes:
    if change is None:
        return WORKBOOK.read_bytes()
    book = openpyxl.load_workbook(WORKBOOK)
    change(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def column(book: Workbook, tab: str, name: str) -> int:
    headers = [str(c.value or "").strip() for c in book[tab][2]]
    return headers.index(name) + 1


def tab(report: dict[str, Any], name: str) -> dict[str, Any]:
    found: dict[str, Any] = next(t for t in report["tabs"] if t["name"] == name)
    return found


# --- T01 / T05 / T06: the valid workbook --------------------------------------------------------


def test_e3_3_t01_consolidated_workbook_upserts(h: ImportHarness) -> None:
    res = h.upload(workbook())
    assert res.status_code == 200, res.text
    report = res.json()
    assert report["applied"] is True and report["mode"] == "upsert"
    assert all(t["errors"] == [] for t in report["tabs"])
    assert tab(report, "Filetype")["created"] == 43
    created = sum(t["created"] for t in report["tabs"])
    assert created == len(report["stable_ids"]) and h.counts() == (created, created)
    assert report["sha256"].startswith("sha256:")
    [audit] = h.actions("MASTERS_IMPORTED")
    assert audit.after is not None
    assert audit.after["sha256"] == report["sha256"] and audit.after["applied"] is True
    assert audit.after["counts"]["Filetype"]["created"] == 43
    assert len(h.actions("MASTER_ENTRY_CREATED")) == created


def test_e3_3_t05_validate_only_writes_nothing(h: ImportHarness) -> None:
    res = h.upload(workbook(), mode="validate-only")
    assert res.status_code == 200, res.text
    report = res.json()
    assert report["applied"] is False and tab(report, "Filetype")["created"] == 43
    assert h.counts() == (0, 0)


def test_e3_3_t06_reimport_is_all_unchanged(h: ImportHarness) -> None:
    assert h.upload(workbook()).status_code == 200
    before = h.counts()
    report = h.upload(workbook()).json()
    assert all(t["created"] == t["updated"] == 0 for t in report["tabs"])
    assert tab(report, "Filetype")["unchanged"] == 43
    assert h.counts() == before


def test_e3_3_changed_row_is_a_new_version(h: ImportHarness) -> None:
    assert h.upload(workbook()).status_code == 200
    entries, versions = h.counts()

    def edit(book: Workbook) -> None:
        cell = book["Rules"].cell(5, column(book, "Rules", "Components (BOM)"))
        cell.value = f"{cell.value or ''} (revised)"

    report = h.upload(workbook(edit)).json()
    assert tab(report, "Rules")["updated"] == 1
    assert h.counts() == (entries, versions + 1)


def test_e3_3_rows_dropped_from_the_workbook_are_kept_and_counted(h: ImportHarness) -> None:
    assert h.upload(workbook()).status_code == 200
    entries, _ = h.counts()

    def drop(book: Workbook) -> None:
        ws = book["Rules"]
        for c in ws[5][1:]:
            c.value = None  # label only = a blank template row

    report = h.upload(workbook(drop)).json()
    assert tab(report, "Rules")["missing"] == 1 and h.counts()[0] == entries


def reorder_rag(rename: dict[str, str] | None = None) -> Callable[[Workbook], None]:
    """Skill #1's RAG Reference with its keys reversed (and optionally renamed)."""

    def change(book: Workbook) -> None:
        cell = book["Skills"].cell(4, column(book, "Skills", "RAG Reference"))
        rag = json.loads(str(cell.value))
        items = [((rename or {}).get(k, k), v) for k, v in reversed(rag.items())]
        cell.value = json.dumps(dict(items))

    return change


def test_e3_3_rag_key_order_is_not_a_change(h: ImportHarness) -> None:
    assert h.upload(workbook()).status_code == 200
    report = h.upload(workbook(reorder_rag())).json()
    assert tab(report, "Skills")["updated"] == 0 and tab(report, "Skills")["errors"] == []


def test_e3_3_renamed_rag_key_is_a_row_error(h: ImportHarness) -> None:
    report = h.upload(workbook(reorder_rag({"ID Field": "Id"}))).json()
    assert report["applied"] is False
    assert [(e["row"], e["column"]) for e in tab(report, "Skills")["errors"]] == [
        (4, "RAG Reference")
    ]


# --- T02 / T03 / T04 / T09: structure, rows and the frozen Filetype -----------------------------


def test_e3_3_t02_renamed_rules_column_is_422_with_a_diff(h: ImportHarness) -> None:
    def rename(book: Workbook) -> None:
        book["Rules"].cell(2, column(book, "Rules", "Components (BOM)")).value = "BOM"

    body = problem(h.upload(workbook(rename)), 422, "STRUCTURE_MISMATCH")
    assert body["errors"][0]["diff"] == {
        "tab": "Rules",
        "expected": "Components (BOM)",
        "actual": "BOM",
        "kind": "renamed",
    }
    assert h.counts() == (0, 0)


def test_e3_3_t03_added_agents_column_is_422_added(h: ImportHarness) -> None:
    def add(book: Workbook) -> None:
        ws = book["Agents"]
        ws.cell(2, column(book, "Agents", "Misc") + 1).value = "Priority"

    body = problem(h.upload(workbook(add)), 422, "STRUCTURE_MISMATCH")
    assert [e["diff"]["kind"] for e in body["errors"]] == ["added"]
    assert body["errors"][0]["diff"]["actual"] == "Priority"
    assert h.counts() == (0, 0)


def test_e3_3_t04_bad_rag_json_is_a_row_error_and_nothing_is_written(h: ImportHarness) -> None:
    def bad(book: Workbook) -> None:
        book["Skills"].cell(5, column(book, "Skills", "RAG Reference")).value = '{"source": '

    res = h.upload(workbook(bad))
    assert res.status_code == 200, res.text
    report = res.json()
    assert report["applied"] is False
    errors = tab(report, "Skills")["errors"]
    assert {(e["tab"], e["row"], e["column"]) for e in errors} == {("Skills", 5, "RAG Reference")}
    assert h.counts() == (0, 0)


def test_e3_3_t09_changed_filetype_tier_is_403(h: ImportHarness) -> None:
    def tier(book: Workbook) -> None:
        cell = book["Filetype"].cell(10, column(book, "Filetype", "Tier"))  # Filetype #7
        cell.value = 3 if cell.value != 3 else 2

    body = problem(h.upload(workbook(tier)), 403, "FILETYPE_FROZEN")
    assert [e["path"] for e in body["errors"]] == ["Filetype!row10:colA"]
    assert h.counts() == (0, 0)


# --- T07 / T08 / T10 / EICAR / size: the file itself ---------------------------------------------


def rezip(extra: dict[str, bytes], content_types: Callable[[str], str] = lambda x: x) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(workbook())) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "[Content_Types].xml":
                data = content_types(data.decode()).encode()
            dst.writestr(info, data)
        for name, data in extra.items():
            dst.writestr(name, data)
    return out.getvalue()


def test_e3_3_t07_xlsm_renamed_to_xlsx_is_415(h: ImportHarness) -> None:
    def macro(xml: str) -> str:
        return xml.replace(
            "spreadsheetml.sheet.main+xml", "../vnd.ms-excel.sheet.macroEnabled.main+xml"
        )

    problem(h.upload(rezip({"xl/vbaProject.bin": b"\0" * 64}, macro)), 415, "UNSUPPORTED_TYPE")
    problem(h.upload(rezip({}, macro)), 415, "UNSUPPORTED_TYPE")  # macro-enabled without the bin
    assert h.counts() == (0, 0)


def test_e3_3_t08_zip_bomb_is_422_before_openpyxl(h: ImportHarness) -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        with z.open("xl/worksheets/sheet1.xml", "w") as member:
            block = b"\0" * (1024 * 1024)
            for _ in range(101):
                member.write(block)
    with mock.patch("sbam_api.masters.importer.openpyxl.load_workbook") as load:
        problem(h.upload(out.getvalue()), 422, "ARCHIVE_LIMIT")
    load.assert_not_called()


def test_e3_3_t10_clamav_down_is_503_and_nothing_is_imported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    with FakeClamd(behaviour="drop") as down:
        h = harness(tmp_path, monkeypatch, down)
        problem(h.upload(workbook()), 503, "SCAN_UNAVAILABLE")
    assert h.counts() == (0, 0) and h.actions("MASTERS_IMPORTED") == []


ZERO = "00000000-0000-0000-0000-000000000000"  # the tenant-wide project id in storage keys


def quarantine_key(data: bytes) -> str:
    return f"{TENANT_A}/{ZERO}/quarantine/{hashlib.sha256(data).hexdigest()}"


def storage_root(h: ImportHarness) -> Path:
    root: Path = h.client.app.state.settings.storage_root  # type: ignore[attr-defined]
    return root


def test_e3_3_eicar_is_422_malware_quarantined_and_audited(
    h: ImportHarness, clamd: FakeClamd
) -> None:
    problem(h.upload(eicar_bytes()), 422, "MALWARE")
    [audit] = h.actions("MASTERS_IMPORT_REJECTED")
    assert audit.after is not None and audit.after["signature"] == "Win.Test.EICAR_HDB-1"
    assert audit.after["quarantine_key"] == quarantine_key(eicar_bytes())
    assert clamd.scans == 1 and h.counts() == (0, 0)


def test_e3_3_infected_workbook_is_quarantined_and_never_parsed(
    h: ImportHarness, clamd: FakeClamd
) -> None:
    clamd.behaviour = "infected"  # a real workbook, so the stored bytes can be read on any OS
    data = workbook()
    with mock.patch("sbam_api.masters.importer.parse_workbook") as parse:
        problem(h.upload(data), 422, "MALWARE")
    parse.assert_not_called()
    assert (storage_root(h) / quarantine_key(data)).read_bytes() == data
    [audit] = h.actions("MASTERS_IMPORT_REJECTED")
    assert audit.after is not None and audit.after["quarantine_key"] == quarantine_key(data)
    assert h.counts() == (0, 0)


def test_e3_3_not_a_zip_is_415(h: ImportHarness) -> None:
    problem(h.upload(b"just text"), 415, "UNSUPPORTED_TYPE")


def test_e3_3_other_extension_is_415(h: ImportHarness) -> None:
    problem(h.upload(workbook(), filename="masters.xls"), 415, "UNSUPPORTED_TYPE")


def test_e3_3_over_15_mb_is_413_and_never_scanned(h: ImportHarness, clamd: FakeClamd) -> None:
    problem(h.upload(b"\0" * (15 * 1024 * 1024 + 1)), 413, "FILE_TOO_LARGE")
    problem(h.upload(b"\0" * (16 * 1024 * 1024 + 1)), 413, "PAYLOAD_TOO_LARGE")
    assert clamd.scans == 0


# --- scope ----------------------------------------------------------------------------------


def test_e3_3_project_scope(h: ImportHarness) -> None:
    project = h.create()
    report = h.upload(workbook(), project_id=project["id"]).json()
    assert report["applied"] is True
    with h.sessions() as s:
        bound = s.scalars(select(MasterEntry.project_id).distinct()).all()
    assert [str(p) for p in bound] == [project["id"]]


def test_e3_3_another_tenants_project_is_404(h: ImportHarness) -> None:
    other = h.call("POST", "/admin/projects", h.admin_b(), json={"name": "B project"}).json()
    problem(h.upload(workbook(), project_id=other["id"]), 404, "NOT_FOUND")
    assert h.counts() == (0, 0)


def test_e3_3_mode_must_be_known(h: ImportHarness) -> None:
    problem(h.upload(workbook(), mode="replace"), 422, "VALIDATION_ERROR")


def test_e3_3_concurrent_edit_during_apply_is_409(h: ImportHarness) -> None:
    assert h.upload(workbook()).status_code == 200

    def edit(book: Workbook) -> None:
        cell = book["Rules"].cell(5, column(book, "Rules", "Components (BOM)"))
        cell.value = f"{cell.value or ''} (revised)"

    # another edit bumps the row between our read and _apply's flush
    with mock.patch("sbam_api.masters.importer._new_version", side_effect=StaleRowVersion()):
        res = h.upload(workbook(edit))
    problem(res, 409, "VERSION_CONFLICT")


# --- A-19: validate before materialising rows -----------------------------------------------------


def _refusal(data: bytes) -> tuple[Any, float]:
    """What `_read` raises for this workbook, and how long it took."""
    import time

    from sbam_api.core.conventions import ApiError
    from sbam_api.masters.importer import _read

    started = time.monotonic()
    with pytest.raises(ApiError) as caught:
        _read("masters.xlsx", data)
    return caught.value, time.monotonic() - started


def test_a19_t01_a_huge_tab_is_refused_from_the_zip_directory_in_under_a_second() -> None:
    """The audit's N-03: a tab with 400k rows (13 MB zipped) took 11 s and 250 MB to reject."""
    import random

    book = Workbook(write_only=True)
    for name in ("README", "Organization", "Rules", "Agents", "Skills", "Filetype"):
        book.create_sheet(name).append(["x"])
    bulk = book.create_sheet("Bulk")
    rng = random.Random(1)  # noqa: S311 - test data
    for n in range(300_000):
        bulk.append([f"{rng.getrandbits(60):x}", f"{rng.getrandbits(60):x}", n])
    out = io.BytesIO()
    book.save(out)
    error, seconds = _refusal(out.getvalue())
    assert error.code == "VALIDATION_ERROR" and error.status == 422
    assert seconds < 1


def test_a19_t02_a_template_tab_over_the_row_cap_is_422_and_stops_streaming(
    h: ImportHarness,
) -> None:
    from sbam_api.masters.importer import MAX_DATA_ROWS

    def grow(book: Workbook) -> None:
        sheet = book["Skills"]
        for n in range(MAX_DATA_ROWS + 50):
            sheet.append([f"filler {n}"])

    body = problem(h.upload(workbook(grow)), 422, "VALIDATION_ERROR")
    assert body["errors"][0]["path"] == "Skills!rows"
    assert h.counts() == (0, 0)  # nothing written


def test_a19_t03_the_cap_counts_filled_rows_not_formatted_empty_ones(h: ImportHarness) -> None:
    from sbam_api.masters.importer import MAX_DATA_ROWS

    def pad(book: Workbook) -> None:
        sheet = book["Skills"]
        sheet.cell(row=MAX_DATA_ROWS + 100, column=1).value = None  # a styled, empty used range
        sheet.row_dimensions[MAX_DATA_ROWS + 100].height = 30

    res = h.upload(workbook(pad), mode="validate-only")
    assert res.status_code == 200, res.text


def test_a19_t04_a_wrong_tab_set_still_reports_the_same_structure_error(h: ImportHarness) -> None:
    def drop(book: Workbook) -> None:
        del book["Rules"]

    body = problem(h.upload(workbook(drop)), 422, "STRUCTURE_MISMATCH")
    assert body["errors"][0]["path"] == "workbook!tabs"
    assert "missing ['Rules']" in body["errors"][0]["message"]
