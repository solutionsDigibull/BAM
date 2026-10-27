"""Masters workbook export and template (E3.4): GET /admin/masters/export and /template.

The round trip runs on the MSC reference workbook: import it, export, import the export again.
"""

import io
import uuid
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from clamd_fake import FakeClamd
from fixtures import fixture_path
from sbam_api.masters.models import MasterEntry, MasterEntryVersion
from sbam_contracts.masters import cell_text
from sbam_contracts.masters_workbook import HEADER_ROWS, TAB_ORDER
from sqlalchemy import select
from test_auth import FDE_A1, problem
from test_masters_import import ImportHarness, harness, tab, workbook

REFERENCE = fixture_path("masters", "BuLLM - Admin Masters (One Sheet) - Consolidated.xlsx")
TENANT_TABS = ("Organization", "Rules", "Agents", "Skills")


@pytest.fixture
def h(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[ImportHarness]:
    with FakeClamd() as clamd:  # the round trip imports, and every import is scanned
        yield harness(tmp_path, monkeypatch, clamd)


def export(h: ImportHarness, token: str | None = None, **params: str) -> Any:
    return h.call("GET", "/admin/masters/export", token or h.admin_a(), params=params)


def book(res: Any) -> openpyxl.Workbook:
    assert res.status_code == 200, res.text
    return openpyxl.load_workbook(io.BytesIO(res.content))


def data_rows(sheet: Any) -> list[list[str]]:
    """Entry rows: a label and at least one value (label-only rows are blank template rows)."""
    rows = sheet.iter_rows(min_row=HEADER_ROWS + 1, values_only=True)
    texts = [[cell_text(v) for v in r] for r in rows]
    return [r for r in texts if any(r[1:])]


def header_rows(sheet: Any) -> list[list[Any]]:
    return [list(r) for r in sheet.iter_rows(max_row=HEADER_ROWS, values_only=True)]


def payloads(h: ImportHarness) -> dict[str, Any]:
    with h.sessions() as s:
        rows = s.execute(
            select(MasterEntry.stable_id, MasterEntryVersion.payload).join(
                MasterEntryVersion,
                (MasterEntryVersion.entry_id == MasterEntry.id)
                & (MasterEntryVersion.version == MasterEntry.current_version),
            )
        ).all()
    return {stable_id: payload for stable_id, payload in rows}


# --- T01: import -> export -> import ------------------------------------------------------------


def test_e3_4_t01_round_trip_is_cell_identical(h: ImportHarness) -> None:
    assert h.upload(workbook(), mode="upsert").json()["applied"] is True
    stored = payloads(h)
    res = export(h)
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    report = h.upload(res.content, mode="validate-only").json()
    assert all(t["errors"] == [] for t in report["tabs"]), report
    assert all(t["created"] == t["updated"] == t["missing"] == 0 for t in report["tabs"]), report
    assert payloads(h) == stored

    exported, original = book(res), openpyxl.load_workbook(REFERENCE)
    assert exported.sheetnames == list(TAB_ORDER)  # no Tooling entries, no Tooling tab
    for name in TAB_ORDER:
        assert header_rows(exported[name]) == header_rows(original[name]), name
    for name in TENANT_TABS:
        want, got = data_rows(original[name]), data_rows(exported[name])
        assert len(got) == len(want), name
        for w, g in zip(want, got, strict=True):
            assert g[0] == w[0].strip()  # the label, without the reference's trailing spaces
            texts = [(a, b) for a, b in zip(w[1:], g[1:], strict=False) if not a.startswith("{")]
            assert all(a == b for a, b in texts), (
                name,
                w[0],
            )  # RAG JSON: same payload, above


def test_e3_4_export_is_audited_and_named_after_the_tenant(h: ImportHarness) -> None:
    report = h.upload(workbook(), mode="upsert").json()
    res = export(h)
    disposition = res.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="BAM_masters_Tenant-A_')
    assert ".xlsx\"; filename*=UTF-8''BAM_masters_Tenant-A_" in disposition  # RFC 6266, E15.7
    assert disposition.endswith(".xlsx") and res.headers["cache-control"] == "no-store"
    [audit] = h.actions("MASTERS_EXPORTED")
    assert (
        audit.after is not None
        and audit.after["Rules"] == tab(report, "Rules")["created"]
        and audit.after["project_id"] is None
    )


def test_e3_4_tenant_with_only_filetype_rows(h: ImportHarness) -> None:
    exported = book(export(h))
    assert all(data_rows(exported[name]) == [] for name in TENANT_TABS)
    filetype = data_rows(exported["Filetype"])
    assert [r[0].strip() for r in filetype[:43]] == [f"Filetype #{n}" for n in range(1, 44)]


# --- T02: template ------------------------------------------------------------------------------


def test_e3_4_t02_template(h: ImportHarness) -> None:
    res = h.call("GET", "/admin/masters/template", h.admin_a())
    template, original = book(res), openpyxl.load_workbook(REFERENCE)
    assert template.sheetnames == list(TAB_ORDER)
    assert [list(r) for r in template["README"].values] == [
        list(r) for r in original["README"].values
    ]
    for name in TAB_ORDER[1:]:
        assert header_rows(template[name]) == header_rows(original[name]), name
        assert template[name].freeze_panes == "A4"
    for name in TENANT_TABS:
        assert all(
            v is None
            for v in next(template[name].iter_rows(min_row=4, max_row=4, values_only=True), ())
        )
        assert data_rows(template[name]) == []
    report = h.upload(res.content, mode="validate-only").json()
    assert all(t["errors"] == [] for t in report["tabs"])
    assert all(
        tab(report, name)["created"] == 0 for name in TENANT_TABS
    )  # 43 frozen Filetype rows only
    assert h.counts() == (0, 0)


# --- T03 / T04: scope ---------------------------------------------------------------------------


def test_e3_4_t03_project_export_has_tenant_and_that_projects_entries(
    h: ImportHarness,
) -> None:
    tenant_rules = tab(h.upload(workbook(), mode="upsert").json(), "Rules")["created"]
    mine, other = h.create("Mine"), h.create("Other")
    for project, rule in ((mine, "RULE-MINE"), (other, "RULE-OTHER")):
        payload = {
            "domain_rule": f"{rule}: project rule. Policy: x",
            "misc": "Status: Proposed",
        }
        body = {"payload": payload, "project_id": project["id"]}
        assert h.call("POST", "/admin/masters/RULES", h.admin_a(), json=body).status_code == 201
    rules = [r[1] for r in data_rows(book(export(h, project_id=mine["id"]))["Rules"])]
    assert any(r.startswith("RULE-MINE") for r in rules)
    assert not any(r.startswith("RULE-OTHER") for r in rules)
    assert len(rules) == tenant_rules + 1  # the tenant's rules + this project's one
    assert len(data_rows(book(export(h))["Rules"])) == tenant_rules  # tenant-wide: no project rows


def test_e3_4_t04_another_tenants_project_is_404(h: ImportHarness) -> None:
    project = h.create()
    problem(export(h, h.admin_b(), project_id=project["id"]), 404, "NOT_FOUND")
    problem(export(h, project_id=str(uuid.uuid4())), 404, "NOT_FOUND")
    assert h.actions("MASTERS_EXPORTED") == []


def test_e3_4_another_tenant_never_sees_my_entries(h: ImportHarness) -> None:
    h.upload(workbook(), mode="upsert")
    exported = book(export(h, h.admin_b()))
    assert all(data_rows(exported[name]) == [] for name in TENANT_TABS)


# --- T05: formula text --------------------------------------------------------------------------


@pytest.mark.parametrize("text", ['=HYPERLINK("http://x")', "+1+1", "-2", "@SUM(A1)"])
def test_e3_4_t05_formula_text_is_a_string_cell(h: ImportHarness, text: str) -> None:
    payload = {
        "domain_rule": "RULE-FX: formulas. Policy: none",
        "pcb_panelization": text,
        "misc": "Status: Proposed",
    }
    assert (
        h.call("POST", "/admin/masters/RULES", h.admin_a(), json={"payload": payload}).status_code
        == 201
    )
    res = export(h)
    with zipfile.ZipFile(io.BytesIO(res.content)) as archive:
        sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/sheet")]
        assert not any(b"<f>" in archive.read(n) or b"<f " in archive.read(n) for n in sheets)
    cell = book(res)["Rules"].cell(4, 3)
    assert cell.value == text and cell.data_type == "s"
    report = h.upload(res.content, mode="validate-only").json()
    assert tab(report, "Rules")["unchanged"] == 1 and tab(report, "Rules")["errors"] == []


# --- T06 and the rest ---------------------------------------------------------------------------


def test_e3_4_t06_fde_cannot_export(h: ImportHarness) -> None:
    fde = h.realm.token(FDE_A1)
    problem(h.call("GET", "/admin/masters/export", fde), 403, "WRONG_PORTAL")
    problem(h.call("GET", "/admin/masters/template", fde), 403, "WRONG_PORTAL")


def test_e3_4_too_long_cell_is_422_naming_the_row(h: ImportHarness) -> None:
    long = "x" * 5_000
    rag = {
        k: long
        for k in (
            "Example",
            "Links / Path",
            "Entity / Node Type",
            "Security / Access Level",
            "Refresh Frequency",
            "ID Field",
            "Parent Entity",
        )
    }
    payload = {
        "domain_rule": "RULE-BIG: big. Policy: x",
        "misc": "Status: Proposed",
        "rag_reference": rag,
    }
    res = h.call("POST", "/admin/masters/RULES", h.admin_a(), json={"payload": payload})
    assert res.status_code == 201, res.text
    body = problem(export(h), 422, "VALIDATION_ERROR")
    assert body["errors"][0]["path"] == "Rules!row4:col12"  # the RAG Reference column
    assert body["errors"][0]["message"].startswith(f"{res.json()['row_label']} has a value")


def test_e3_4_tooling_tab_only_when_there_are_tooling_entries(h: ImportHarness) -> None:
    payload = {"category": "MCP", "examples": "WebMCP", "record_status": "Proposed"}
    body = {"payload": payload, "row_label": "TOOL-MFG-01"}
    res = h.call("POST", "/admin/masters/TOOLING", h.admin_a(), json=body)
    assert res.status_code == 201, res.text
    exported = book(export(h))
    assert exported.sheetnames == [
        "README",
        "Organization",
        "Rules",
        "Agents",
        "Skills",
        "Tooling",
        "Filetype",
    ]
    assert data_rows(exported["Tooling"])[0][:3] == ["TOOL-MFG-01", "MCP", "WebMCP"]
    report = h.upload(_bytes(exported), mode="validate-only").json()
    assert tab(report, "Tooling")["unchanged"] == 1 and tab(report, "Tooling")["errors"] == []


def _bytes(workbook_: openpyxl.Workbook) -> bytes:
    out = io.BytesIO()
    workbook_.save(out)
    return out.getvalue()
