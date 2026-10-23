# ruff: noqa: F811  (the fixtures `h` and `root` come from test_parsing.py by import)
"""E4.5: PDF, DOCX and XLSX through the `parse_document` step (storage, state machine, report).

The loaders alone are in test_loaders_office.py; the harness is the one of test_parsing.py.
"""

import io
import socket
from typing import Any

import pytest
from fixtures import fixture_path
from openpyxl import Workbook
from pdfs import encrypted, make_pdf
from sbam_contracts.api import ErrorCode
from sbam_workers.ingest import parsing
from test_loaders_office import sample_docx, sheet_of, xlsx_bytes
from test_parsing import Harness, derived, h, root  # noqa: F401 - their fixtures

PDF_ROW = {"detected": "pdf", "filetype_row_label": "Filetype #21"}
DOCX_ROW = {"detected": "docx", "filetype_row_label": "Filetype #19"}
XLSX_ROW = {"detected": "xlsx", "filetype_row_label": "Filetype #23"}


def two_sheets() -> bytes:
    def build(book: Workbook) -> None:
        first = sheet_of(book)
        first.title = "Sheet1"
        first.append(["MPN", "Qty"])
        first.append(["STM32", 12])
        first["B4"] = 99
        book.create_sheet("Prices").append(["Item", "EUR"])

    return xlsx_bytes(build)


# --- parsed ---------------------------------------------------------------------------------------


def test_e4_5_t01_a_ten_page_pdf_reaches_parsed_with_a_block_for_every_page(h: Harness) -> None:
    ref = h.add(
        "offer.pdf",
        make_pdf([f"Line {n} with the parts and prices" for n in range(1, 11)]),
        **PDF_ROW,
    )
    outcome = parsing.parse_version(h.deps, ref)
    assert outcome.parsed and outcome.report is not None and outcome.report.pages == 10
    assert h.version(ref).state == "PARSED"
    md, doc = derived(h, ref)
    assert [b.page for b in doc.blocks] == list(range(1, 11))
    assert md.decode().startswith("<!-- page 1 -->\n\nLine 1 with the parts and prices\n")
    assert doc.loader == "Filetype #21"


def test_e4_5_t02_a_two_sheet_workbook_reaches_parsed_with_sheet_and_cell_provenance(
    h: Harness,
) -> None:
    ref = h.add("rates.xlsx", two_sheets(), **XLSX_ROW)
    outcome = parsing.parse_version(h.deps, ref)
    assert outcome.parsed
    md, doc = derived(h, ref)
    found = {(b.sheet, b.cell): b.text for b in doc.blocks if b.cell}
    assert found[("Sheet1", "B4")] == "99" and found[("Sheet1", "A2")] == "STM32"
    assert [t.sheet for t in doc.tables] == ["Sheet1", "Prices"]
    assert md.decode().startswith("## Sheet1\n\n| MPN | Qty |")
    version = h.version(ref)
    assert version.parse_report is not None and version.parse_report["sheets"] == 2
    assert h.audit().count("DOCUMENT_PARSED") == 1


def test_e4_5_t04_a_docx_keeps_its_headings_and_gets_a_table(h: Harness) -> None:
    ref = h.add("terms.docx", sample_docx(), **DOCX_ROW)
    assert parsing.parse_version(h.deps, ref).parsed
    md, doc = derived(h, ref)
    assert md.decode().startswith("# Quote terms\n\n# Scope\n")
    assert "## Prices\n" in md.decode()
    assert doc.tables[0].header == ["MPN", "Qty"] and doc.tables[0].rows == [["STM32", "12"]]


def test_e4_5_pages_that_need_ocr_are_kept_in_the_report_for_the_ocr_step(h: Harness) -> None:
    mixed = h.add(
        "scan.pdf", fixture_path("documents", "scanned_mixed.pdf").read_bytes(), **PDF_ROW
    )
    outcome = parsing.parse_version(h.deps, mixed)
    # Not parsed yet: the version waits in PARSING for the OCR step (E4.8)
    assert (outcome.parsed, outcome.code) == (False, None) and outcome.report is not None
    assert outcome.report.ocr_needed_pages == [2]
    assert h.version(mixed).state == "PARSING" and h.version(mixed).parse_report is not None
    assert h.version(mixed).parse_report["ocr_needed_pages"] == [2]  # type: ignore[index]
    full = h.add("all.pdf", fixture_path("documents", "scanned_full.pdf").read_bytes(), **PDF_ROW)
    waiting = parsing.parse_version(h.deps, full)
    assert not waiting.parsed and waiting.report is not None
    assert waiting.report.ocr_needed_pages == [1, 2]
    md, doc = derived(h, full)  # the text pass is stored: nothing to keep yet, the OCR adds to it
    assert md == b"" and doc.blocks == []


def test_e4_5_a_workbook_is_written_out_block_by_block_and_every_block_is_counted(
    h: Harness,
) -> None:
    def build(book: Workbook) -> None:
        for n in range(1, 3001):
            sheet_of(book).append([n, f"row {n}", None, "x"])

    ref = h.add("big.xlsx", xlsx_bytes(build), **XLSX_ROW)
    assert parsing.parse_version(h.deps, ref).parsed
    _, doc = derived(h, ref)
    cell_blocks = [b for b in doc.blocks if b.type.value == "cell"]
    assert len(cell_blocks) == 3000 * 3  # the empty column C has no block
    assert len(doc.tables[0].rows) == 2999
    with h.sessions() as s:
        from sbam_api.audit.models import AuditLog
        from sqlalchemy import select

        audit = s.scalars(select(AuditLog).where(AuditLog.action == "DOCUMENT_PARSED")).one()
    assert audit.after is not None and audit.after["blocks"] == len(doc.blocks)


def test_e4_5_the_activity_beats_its_heart_while_a_long_workbook_is_parsed(h: Harness) -> None:
    def build(book: Workbook) -> None:
        for n in range(2500):
            sheet_of(book).append([n, "x"])

    ref = h.add("long.xlsx", xlsx_bytes(build), **XLSX_ROW)
    beats: list[int] = []
    parsing.parse_version(h.deps, ref, lambda: beats.append(1))
    assert len(beats) >= 2


# --- failed ---------------------------------------------------------------------------------------


def test_e4_5_t05_an_encrypted_pdf_is_parse_failed_with_its_code_and_no_retry(h: Harness) -> None:
    locked = encrypted(make_pdf(["secret"]), user="pw", owner="pw")
    ref = h.add("locked.pdf", locked, **PDF_ROW)
    outcome = parsing.parse_version(
        h.deps, ref
    )  # returned, not raised: nothing for Temporal to retry
    assert (outcome.parsed, outcome.code) == (False, ErrorCode.ENCRYPTED)
    version = h.version(ref)
    assert (version.state, version.parse_code) == ("PARSE_FAILED", "ENCRYPTED")
    assert version.derived_json_key is None
    assert "document_version.parse_failed" in h.outbox()


@pytest.mark.parametrize(
    ("name", "data", "row", "code"),
    [
        ("empty.docx", None, DOCX_ROW, ErrorCode.EMPTY),
        ("broken.docx", b"PK\x03\x04 broken", DOCX_ROW, ErrorCode.CORRUPT),
        ("broken.pdf", b"%PDF-1.4 broken", PDF_ROW, ErrorCode.CORRUPT),
        ("broken.xlsx", b"PK\x03\x04 broken", XLSX_ROW, ErrorCode.CORRUPT),
        ("empty.xlsx", "empty", XLSX_ROW, ErrorCode.EMPTY),
    ],
    ids=["empty-docx", "corrupt-docx", "corrupt-pdf", "corrupt-xlsx", "empty-xlsx"],
)
def test_e4_5_a_file_that_cannot_be_read_is_parse_failed_with_its_code(
    h: Harness, name: str, data: Any, row: dict[str, str], code: ErrorCode
) -> None:
    if data is None:
        from docx import Document

        out = io.BytesIO()
        Document().save(out)
        data = out.getvalue()
    elif data == "empty":
        data = xlsx_bytes(lambda book: None)
    ref = h.add(name, data, **row)
    outcome = parsing.parse_version(h.deps, ref)
    assert (outcome.parsed, outcome.code) == (False, code)
    assert (h.version(ref).state, h.version(ref).parse_code) == ("PARSE_FAILED", code.value)


def test_e4_5_a_derivative_over_the_size_limit_is_parse_failed_not_retried(
    h: Harness, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(parsing, "DERIVED_MAX_BYTES", 200)
    ref = h.add("rates.xlsx", two_sheets(), **XLSX_ROW)
    outcome = parsing.parse_version(h.deps, ref)
    assert (outcome.parsed, outcome.code) == (False, ErrorCode.PARSE_ERROR)
    assert h.version(ref).state == "PARSE_FAILED"


def test_e4_5_t07_a_workbook_with_a_formula_is_parsed_without_touching_the_network(
    h: Harness, monkeypatch: pytest.MonkeyPatch
) -> None:
    def build(book: Workbook) -> None:
        sheet_of(book)["A1"] = '=WEBSERVICE("http://evil.example/x")'

    def no_network(*args: Any, **kwargs: Any) -> None:
        raise AssertionError("the parser touched the network")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    monkeypatch.setattr(socket, "getaddrinfo", no_network)
    ref = h.add("sneaky.xlsx", xlsx_bytes(build), **XLSX_ROW)
    assert parsing.parse_version(h.deps, ref).parsed
    _, doc = derived(h, ref)
    assert [b.text for b in doc.blocks if b.cell] == ['=WEBSERVICE("http://evil.example/x")']
