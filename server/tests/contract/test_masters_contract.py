"""E0.4 contract tests: the consolidated masters workbook against the BAM 0.2 contract.

The workbook is read straight from docs/reference (no copy; the fixture pack is issue E0.8).
Cells are read as written (`data_only=False`), so nothing is ever evaluated.
"""

import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from sbam_contracts.api import ErrorCode
from sbam_contracts.masters import (
    AgentsPayload,
    FiletypePayload,
    MasterStatus,
    MasterType,
    OrganizationPayload,
    RagReference,
    ReadmeTab,
    RulesPayload,
    SkillsPayload,
    SourceKind,
    ToolingPayload,
    schemas,
)
from sbam_contracts.masters_workbook import (
    COLUMNS,
    FILETYPE_COUNT,
    MAX_WORKBOOK_BYTES,
    PAYLOAD_MODELS,
    TAB_ORDER,
    WorkbookReport,
    check_workbook_container,
    parse_workbook,
    validate_structure,
)

REPO = Path(__file__).resolve().parents[3]
WORKBOOK = REPO / "docs/reference/BuLLM - Admin Masters (One Sheet) - Consolidated.xlsx"
SCHEMAS = REPO / "server/contracts/schemas/masters"

Sheets = dict[str, list[list[Any]]]


def load_sheets() -> Sheets:
    book = openpyxl.load_workbook(WORKBOOK, data_only=False, read_only=True)
    try:
        return {
            ws.title: [list(row) for row in ws.iter_rows(values_only=True)]
            for ws in book.worksheets
        }
    finally:
        book.close()


@pytest.fixture(name="sheets")
def fixture_sheets() -> Sheets:
    return load_sheets()


def mutated(sheets: Sheets) -> Sheets:
    return {tab: [list(row) for row in grid] for tab, grid in sheets.items()}


def paths(report: WorkbookReport) -> list[str]:
    return [error.path for error in report.errors]


# --- E0.4-T01: the whole workbook validates ------------------------------------------------------


def test_e0_4_t01_full_workbook_validates(sheets: Sheets) -> None:
    report = parse_workbook(sheets)
    assert report.errors == []
    assert report.code is None


def test_e0_4_t01_row_counts_per_master(sheets: Sheets) -> None:
    rows = parse_workbook(sheets).rows
    counts = {t: sum(1 for r in rows if r.master_type == t) for t in MasterType}
    assert counts == {
        MasterType.ORGANIZATION: 4,  # Unit #5 is a blank template row
        MasterType.RULES: 6,
        MasterType.AGENTS: 5,  # Agent #6 is a blank template row
        MasterType.SKILLS: 8,
        MasterType.TOOLING: 0,  # the BAM 0.2 one-sheet workbook has no Tooling tab
        MasterType.FILETYPE: FILETYPE_COUNT,
    }


def test_e0_4_t01_filetype_has_exactly_43_rows_and_a_note_that_is_not_an_entry(
    sheets: Sheets,
) -> None:
    assert FILETYPE_COUNT == 43
    filetype_rows = [r for r in parse_workbook(sheets).rows if r.master_type == MasterType.FILETYPE]
    assert [r.row_label.strip() for r in filetype_rows] == [f"Filetype #{n}" for n in range(1, 44)]
    last_labels = [row[0] for row in sheets["Filetype"] if row and row[0]]
    assert str(last_labels[-1]).startswith("Note:")


def test_e0_4_t01_stable_ids_and_statuses_for_every_row(sheets: Sheets) -> None:
    by_type: dict[MasterType, list[tuple[str, str, str]]] = {}
    for row in parse_workbook(sheets).rows:
        identity = row.identity
        by_type.setdefault(row.master_type, []).append(
            (identity.stable_id, identity.status, identity.source_kind)
        )
    assert by_type[MasterType.ORGANIZATION] == [
        ("BU-AERO-01", "ILLUSTRATIVE", "BAM_0_2_SAMPLE"),
        ("ORG-EMS", "PROPOSED", "TEMPLATE_ROW"),
        ("ORG-CAB", "PROPOSED", "TEMPLATE_ROW"),
        ("ORG-CORE", "ILLUSTRATIVE", "TEMPLATE_ROW"),
    ]
    assert [(s, st) for s, st, _ in by_type[MasterType.RULES]] == [
        ("Rule #1", "ILLUSTRATIVE"),
        ("RULE-BOM", "PROPOSED"),
        ("RULE-OFFER", "PROPOSED"),
        ("RULE-EXCESS", "PROPOSED"),
        ("RULE-REPORT", "PROPOSED"),
        ("RULE-CORE", "ILLUSTRATIVE"),
    ]
    assert [(s, st) for s, st, _ in by_type[MasterType.AGENTS]] == [
        ("Agent #1", "ILLUSTRATIVE"),
        ("AG-INTAKE", "PROPOSED"),
        ("AG-OFFER", "PROPOSED"),
        ("AG-COST", "PROPOSED"),
        ("AG-QA", "PROPOSED"),
    ]
    assert [(s, st) for s, st, _ in by_type[MasterType.SKILLS]] == [
        ("Skill #1", "ILLUSTRATIVE"),
        ("SK-INTAKE", "PROPOSED"),
        ("SK-BOM", "PROPOSED"),
        ("SK-OFFER", "PROPOSED"),
        ("SK-COST", "PROPOSED"),
        ("SK-ANALYZE", "PROPOSED"),
        ("SK-EXPORT", "PROPOSED"),
        ("SK-CORE", "ILLUSTRATIVE"),
    ]
    assert all(
        st == "PROPOSED" and kind == "TEMPLATE_ROW" for _, st, kind in by_type[MasterType.FILETYPE]
    )


def test_e0_4_t01_sample_rows_are_exactly_row_one_of_each_editable_tab(sheets: Sheets) -> None:
    samples = [
        (r.master_type, r.row_label.strip())
        for r in parse_workbook(sheets).rows
        if r.identity.source_kind == SourceKind.BAM_0_2_SAMPLE
    ]
    assert samples == [
        (MasterType.ORGANIZATION, "Unit #1"),
        (MasterType.RULES, "Rule #1"),
        (MasterType.AGENTS, "Agent #1"),
        (MasterType.SKILLS, "Skill #1"),
    ]


def test_e0_4_t01_rag_references_parse_into_seven_key_objects(sheets: Sheets) -> None:
    rows = parse_workbook(sheets).rows
    refs = [
        r.payload.rag_reference
        for r in rows
        if isinstance(r.payload, RulesPayload | SkillsPayload) and r.payload.rag_reference
    ]
    assert [ref.id_field for ref in refs] == ["DOM-PCBA-01", "PRJ-RFQ-2026-001"]
    assert all(isinstance(ref, RagReference) for ref in refs)


def test_e0_4_t01_no_sbam_metadata_in_payloads_or_columns() -> None:
    metadata = {"stable_id", "status", "source_kind", "row_label", "version", "scope"}
    for model in PAYLOAD_MODELS.values():
        assert metadata.isdisjoint(model.model_fields), model.__name__
    for columns in COLUMNS.values():
        assert metadata.isdisjoint(field for _, field in columns)


def test_e0_4_t01_column_tables_match_the_models_field_for_field() -> None:
    for master_type, model in PAYLOAD_MODELS.items():
        assert [field for _, field in COLUMNS[master_type]] == list(model.model_fields), master_type
    assert [len(COLUMNS[t]) for t in MasterType] == [14, 11, 11, 11, 15, 5]


def test_e0_4_t01_column_names_come_from_the_workbook_row_2(sheets: Sheets) -> None:
    for master_type, tab in (
        (MasterType.ORGANIZATION, "Organization"),
        (MasterType.RULES, "Rules"),
        (MasterType.AGENTS, "Agents"),
        (MasterType.SKILLS, "Skills"),
        (MasterType.FILETYPE, "Filetype"),
    ):
        header = [str(c).strip() for c in sheets[tab][1] if c is not None]
        assert header[1:] == [name for name, _ in COLUMNS[master_type]], tab


# --- Committed JSON Schemas ---------------------------------------------------------------------


@pytest.mark.parametrize("stem", sorted(schemas()))
def test_e0_4_t01_committed_schema_matches_model(stem: str) -> None:
    committed = json.loads((SCHEMAS / f"{stem}.schema.json").read_text(encoding="utf-8"))
    assert (
        committed == schemas()[stem]
    ), "regenerate: cd server && uv run python -m sbam_contracts.masters"


def test_e0_4_t01_schema_files_are_exactly_the_documented_ones() -> None:
    names = {p.name for p in SCHEMAS.glob("*.json")}
    assert names == {
        "readme.schema.json",
        "organization.schema.json",
        "rules.schema.json",
        "agents.schema.json",
        "skills.schema.json",
        "filetype.schema.json",
        "rag_reference.schema.json",
        "tooling.schema.json",  # Manufacturing Masters v0.2 (E3.2)
        # masters engine API (E3.2)
        "master_entry_create.schema.json",
        "master_version_create.schema.json",
        "filetype_info.schema.json",  # E3.7
        "master_entry.schema.json",
        "masters_snapshot.schema.json",
    }


@pytest.mark.parametrize(
    "stem", ["readme", "organization", "rules", "agents", "skills", "tooling", "filetype"]
)
def test_e0_4_t01_schemas_carry_the_structure_version_and_forbid_extras(stem: str) -> None:
    schema = json.loads((SCHEMAS / f"{stem}.schema.json").read_text(encoding="utf-8"))
    assert schema["structure_version"] == "BAM 0.2"
    assert schema["additionalProperties"] is False


def test_e0_4_t01_filetype_schema_is_marked_system_owned() -> None:
    filetype = json.loads((SCHEMAS / "filetype.schema.json").read_text(encoding="utf-8"))
    assert filetype["x-system-owned"] is True
    for stem in ("organization", "rules", "agents", "skills", "tooling"):
        schema = json.loads((SCHEMAS / f"{stem}.schema.json").read_text(encoding="utf-8"))
        assert "x-system-owned" not in schema


def test_e0_4_t05_rag_reference_schema_has_seven_required_keys_and_no_extras() -> None:
    schema = json.loads((SCHEMAS / "rag_reference.schema.json").read_text(encoding="utf-8"))
    assert schema["additionalProperties"] is False
    assert sorted(schema["required"]) == sorted(
        [
            "Example",
            "Links / Path",
            "Entity / Node Type",
            "Security / Access Level",
            "Refresh Frequency",
            "ID Field",
            "Parent Entity",
        ]
    )


def test_e0_4_t01_models_are_the_ones_the_schemas_describe() -> None:
    assert ReadmeTab.model_json_schema().items() <= schemas()["readme"].items()  # + $schema/$id
    assert {
        OrganizationPayload,
        RulesPayload,
        AgentsPayload,
        SkillsPayload,
        ToolingPayload,
        FiletypePayload,
    } == set(PAYLOAD_MODELS.values())


# --- E0.4-T06: structure changes are rejected, never auto-mapped ---------------------------------


def test_e0_4_t06_renamed_column_is_rejected_with_its_path(sheets: Sheets) -> None:
    changed = mutated(sheets)
    misc = [str(c).strip() if c else c for c in changed["Rules"][1]].index("Misc")
    changed["Rules"][1][misc] = "Notes"
    report = parse_workbook(changed)
    assert report.code == ErrorCode.STRUCTURE_MISMATCH
    assert paths(report) == ["Rules!row2:col9"]
    assert report.rows == []  # nothing is imported from a changed structure


def test_e0_4_t06_extra_column_is_rejected_with_its_path(sheets: Sheets) -> None:
    changed = mutated(sheets)
    row = changed["Rules"][1]
    row[:] = [*row[:12], "Priority"]  # a 12th data column after RAG Reference (numbered 0..10)
    report = parse_workbook(changed)
    assert report.code == ErrorCode.STRUCTURE_MISMATCH
    assert paths(report) == ["Rules!row2:col11"]
    assert "unexpected column 'Priority'" in report.errors[0].message


def test_e0_4_t06_removed_column_is_rejected(sheets: Sheets) -> None:
    changed = mutated(sheets)
    changed["Agents"][1] = [c for c in changed["Agents"][1] if str(c).strip() != "Eval"]
    errors = validate_structure(changed)
    assert errors, "a removed column must be reported"
    assert errors[0].path.startswith("Agents!row2:col")


def test_e0_4_t06_reordered_columns_are_rejected(sheets: Sheets) -> None:
    changed = mutated(sheets)
    header = changed["Skills"][1]
    header[2], header[3] = header[3], header[2]
    report = parse_workbook(changed)
    assert report.code == ErrorCode.STRUCTURE_MISMATCH
    assert len(report.errors) == 2


def test_e0_4_t06_renamed_or_missing_or_reordered_tab_is_rejected(sheets: Sheets) -> None:
    renamed = {("Rule" if tab == "Rules" else tab): grid for tab, grid in sheets.items()}
    assert validate_structure(renamed)[0].path == "workbook!tabs"
    missing = {tab: grid for tab, grid in sheets.items() if tab != "Agents"}
    assert "missing ['Agents']" in validate_structure(missing)[0].message
    reordered = {tab: sheets[tab] for tab in ("Organization", "README", *TAB_ORDER[2:])}
    assert "reordered" in validate_structure(reordered)[0].message


def test_e0_4_t06_filetype_with_42_rows_is_rejected(sheets: Sheets) -> None:
    changed = mutated(sheets)
    del changed["Filetype"][45]  # sheet row 46 = Filetype #43
    report = parse_workbook(changed)
    assert report.code == ErrorCode.STRUCTURE_MISMATCH
    assert report.errors[0].path.startswith("Filetype!row")


def test_e0_4_t06_whitespace_only_header_differences_are_not_a_structure_change(
    sheets: Sheets,
) -> None:
    # The shipped workbook has ' Logs' and 'Name ' with stray spaces; they must keep validating.
    changed = mutated(sheets)
    changed["Agents"][1] = [f" {c} " if c else c for c in changed["Agents"][1]]
    assert validate_structure(changed) == []


# --- Row-level errors carry their path ------------------------------------------------------------


def test_e0_4_t07_unknown_organization_type_in_the_workbook_is_reported_with_its_path(
    sheets: Sheets,
) -> None:
    changed = mutated(sheets)
    changed["Organization"][4][2] = "Plant"  # Unit #2, column "Type" (numbered 2)
    report = parse_workbook(changed)
    assert report.code == ErrorCode.VALIDATION_ERROR
    assert paths(report) == ["Organization!row5:col2"]


def test_e0_4_t07_filetype_tier_outside_1_to_3_is_reported(sheets: Sheets) -> None:
    changed = mutated(sheets)
    changed["Filetype"][3][2] = 4.0
    assert paths(parse_workbook(changed)) == ["Filetype!row4:col2"]
    changed["Filetype"][3][2] = 0
    assert paths(parse_workbook(changed)) == ["Filetype!row4:col2"]


def test_e0_4_t04_unknown_status_in_the_workbook_is_reported_at_the_misc_cell(
    sheets: Sheets,
) -> None:
    changed = mutated(sheets)
    changed["Rules"][4][10] = "Source fields needed. Status: Foo"  # Rule #2 Misc (numbered 9)
    report = parse_workbook(changed)
    assert paths(report) == ["Rules!row5:col9"]
    assert report.errors[0].message == "unknown status Foo"


def test_e0_4_t05_broken_rag_cell_in_the_workbook_is_reported(sheets: Sheets) -> None:
    changed = mutated(sheets)
    changed["Rules"][3][11] = "[1, 2]"  # Rule #1 RAG Reference (numbered 10)
    report = parse_workbook(changed)
    assert paths(report) == ["Rules!row4:col10"]
    changed["Rules"][3][11] = json.dumps({"Example": "only one key"})
    assert paths(parse_workbook(changed)) == ["Rules!row4:col10"]


def test_e0_4_t01_duplicate_stable_id_is_reported(sheets: Sheets) -> None:
    changed = mutated(sheets)
    changed["Agents"][5][2] = str(changed["Agents"][4][2])  # Agent #3 now carries AG-INTAKE too
    report = parse_workbook(changed)
    assert any("duplicate stable id AG-INTAKE" in e.message for e in report.errors)


def test_e0_4_t01_bad_row_label_is_reported(sheets: Sheets) -> None:
    changed = mutated(sheets)
    changed["Skills"][4][0] = "Step #2"
    assert paths(parse_workbook(changed)) == ["Skills!row5:colA"]


# --- E0.4-T08: formula injection ------------------------------------------------------------------


def test_e0_4_t08_formula_text_in_the_workbook_stays_literal(sheets: Sheets) -> None:
    changed = mutated(sheets)
    evil = '=HYPERLINK("http://evil")'
    changed["Organization"][4][1] = evil  # Unit #2 Org Name
    report = parse_workbook(changed)
    assert report.errors == []
    unit2 = next(r for r in report.rows if r.identity.stable_id == "ORG-EMS")
    assert isinstance(unit2.payload, OrganizationPayload)
    assert unit2.payload.org_name == evil


def test_e0_4_t08_the_shipped_workbook_holds_no_real_formulas() -> None:
    book = openpyxl.load_workbook(WORKBOOK, data_only=False)
    formulas = [
        c.coordinate
        for ws in book.worksheets
        for r in ws.iter_rows()
        for c in r
        if c.data_type == "f"
    ]
    assert formulas == []


# --- E0.4-T09: only plain .xlsx, at most 15 MB --------------------------------------------------


def test_e0_4_t09_xlsm_is_rejected() -> None:
    errors = check_workbook_container(
        "masters.xlsm", 1_000, ["xl/workbook.xml", "xl/vbaProject.bin"]
    )
    assert {e.message for e in errors} == {
        "only .xlsx workbooks are accepted",
        "workbooks with macros are not accepted",
    }


def test_e0_4_t09_macro_payload_is_rejected_even_with_an_xlsx_name() -> None:
    errors = check_workbook_container("masters.xlsx", 1_000, ["xl/VBAProject.bin"])
    assert [e.message for e in errors] == ["workbooks with macros are not accepted"]


def test_e0_4_t09_other_file_types_are_rejected() -> None:
    assert check_workbook_container("masters.csv", 10, [])[0].path == "file"
    assert check_workbook_container("masters.xls", 10, [])[0].path == "file"


def test_e0_4_t09_size_limit_is_15_mb() -> None:
    assert check_workbook_container("m.xlsx", MAX_WORKBOOK_BYTES, ["xl/workbook.xml"]) == []
    errors = check_workbook_container("m.xlsx", MAX_WORKBOOK_BYTES + 1, ["xl/workbook.xml"])
    assert [e.message for e in errors] == ["workbook is larger than 15 MB"]


def test_e0_4_t09_the_shipped_workbook_passes_the_container_check() -> None:
    import zipfile

    with zipfile.ZipFile(WORKBOOK) as archive:
        names = archive.namelist()
    assert check_workbook_container(WORKBOOK.name, WORKBOOK.stat().st_size, names) == []


def test_e0_4_t01_status_enum_is_the_documented_set() -> None:
    assert [s.value for s in MasterStatus] == ["PROPOSED", "ILLUSTRATIVE", "APPROVED", "RETIRED"]
