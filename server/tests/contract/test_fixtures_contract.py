"""E0.8 contract tests: the ISA14 manifest, masters workbook and SOW fixture are loaded and checked.

Each of the three fixtures is used by at least one test here (E0.8-T01). The ISA14 sample is also
validated against the E0.3 model (`Isa14Manifest`, #20) and the SOW fixtures against `SowDocument`
(E0.5, #22).
"""

import hashlib
import json
import re
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import openpyxl
from fixtures import FIXTURES_DIR, document_version_uuid, fixture_path
from sbam_contracts.masters_workbook import TAB_ORDER, parse_workbook

REPO = Path(__file__).resolve().parents[3]
WORKBOOK_NAME = "BuLLM - Admin Masters (One Sheet) - Consolidated.xlsx"
SOW_SECTIONS = [
    "overview",
    "users_authority",
    "scope",
    "modules",
    "features",
    "data_integrations",
    "rules_constraints",
    "deliverables_acceptance",
    "assumptions_open_points",
]
FEATURE_FIELDS = {
    "id",
    "module_id",
    "name",
    "trigger",
    "action",
    "check",
    "bound",
    "inputs_state",
    "external_data",
    "tools",
    "tables",
    "success_criteria",
    "output",
    "handoff",
    "notifications",
    "citations",
}


def load_sow() -> dict[str, Any]:
    data: dict[str, Any] = json.loads(fixture_path("sow", "sow_rfq_nav9.json").read_text("utf-8"))
    return data


def all_citations(sow: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    found: list[tuple[str, dict[str, Any]]] = []
    for name, section in sow.items():
        items = section if isinstance(section, list) else [section]
        for item in items:
            found.extend((name, c) for c in item.get("citations", []))
    return found


# --- ISA14 sample manifest -----------------------------------------------------------------------


def test_e0_8_t01_isa14_manifest_loads_with_the_five_blocks() -> None:
    manifest = json.loads(fixture_path("isa14", "ms01_1_sample.json").read_text("utf-8"))
    assert list(manifest) == [
        "skill_metadata",
        "isatvon_definition",
        "srim_runtime_definition",
        "contracts",
        "governance",
    ]
    meta = manifest["skill_metadata"]
    assert meta["skill_id"] == "MS01.1_Investor_Auth_Manager"
    assert meta["version"] == "1.4.0"


def test_e0_8_t01_isa14_manifest_validates_against_the_e0_3_model() -> None:
    from sbam_contracts.canonical import canonical_json
    from sbam_contracts.manifest import Isa14Manifest

    manifest = json.loads(fixture_path("isa14", "ms01_1_sample.json").read_text("utf-8"))
    model = Isa14Manifest.model_validate(manifest)
    assert canonical_json(model.dump()) == canonical_json(manifest)


def test_e0_8_t01_isa14_manifest_matches_the_reference_document() -> None:
    import zipfile

    docx = REPO / "docs/reference/ISATVON 1.4 for Skill Generation.docx"
    with zipfile.ZipFile(docx) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    manifest = json.loads(fixture_path("isa14", "ms01_1_sample.json").read_text("utf-8"))
    assert manifest["skill_metadata"]["skill_id"] in xml
    assert (
        escape(manifest["skill_metadata"]["skill_name"]) in xml
    )  # the docx stores XML-escaped text


# --- Masters workbook ----------------------------------------------------------------------------


def test_e0_8_t01_masters_workbook_is_an_unmodified_copy() -> None:
    copy = fixture_path("masters", WORKBOOK_NAME).read_bytes()
    original = (REPO / "docs/reference" / WORKBOOK_NAME).read_bytes()
    assert hashlib.sha256(copy).hexdigest() == hashlib.sha256(original).hexdigest()
    assert copy == original


def test_e0_8_t01_masters_workbook_loads_with_the_six_frozen_tabs() -> None:
    book = openpyxl.load_workbook(fixture_path("masters", WORKBOOK_NAME), read_only=True)
    try:
        assert tuple(book.sheetnames) == TAB_ORDER
    finally:
        book.close()


def test_e0_8_t01_masters_workbook_passes_the_masters_contract() -> None:
    book = openpyxl.load_workbook(fixture_path("masters", WORKBOOK_NAME), data_only=False)
    sheets = {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in book}
    report = parse_workbook(sheets)
    assert report.errors == []
    assert len(report.rows) == 66


# --- SOW fixture ---------------------------------------------------------------------------------


def test_e0_8_t01_sow_fixture_has_the_nine_sections_in_order_and_none_empty() -> None:
    sow = load_sow()
    assert list(sow) == SOW_SECTIONS
    assert all(sow[name] for name in SOW_SECTIONS)


def test_e0_8_t01_every_sow_section_and_feature_is_cited() -> None:
    sow = load_sow()
    for name in SOW_SECTIONS:
        section = sow[name]
        items = section if isinstance(section, list) else [section]
        assert all(item["citations"] for item in items), name


def test_e0_8_t01_sow_citations_point_at_real_fixture_files() -> None:
    files = {
        str(document_version_uuid(path.relative_to(FIXTURES_DIR).as_posix()))
        for path in (FIXTURES_DIR / "documents").iterdir()
    }
    for name in ("sow_rfq_nav9.json", "ms01_1_sow.json"):
        sow = json.loads(fixture_path("sow", name).read_text("utf-8"))
        for section, citation in all_citations(sow):
            assert set(citation) == {"okf_ref_id", "document_version_id", "location"}, section
            assert citation["document_version_id"] in files, (name, section)


def test_e0_8_t01_sow_fixtures_validate_as_sow_document() -> None:
    from sbam_contracts.sow import SowDocument

    for name in ("sow_rfq_nav9.json", "ms01_1_sow.json"):
        SowDocument.model_validate(json.loads(fixture_path("sow", name).read_text("utf-8")))


def test_e0_8_t01_sow_feature_items_have_stable_ids_and_the_full_field_set() -> None:
    sow = load_sow()
    ids = [feature["id"] for feature in sow["features"]]
    assert len(ids) == len(set(ids))
    module_ids = {module["id"] for module in sow["modules"]}
    for feature in sow["features"]:
        assert set(feature) == FEATURE_FIELDS
        assert re.fullmatch(r"F-[A-Z]+-\d+", feature["id"])
        assert feature["module_id"] in module_ids
        assert set(feature["tools"]) == {"action", "check"}
